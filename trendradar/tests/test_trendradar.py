"""Tests: draai met  python -m pytest  in de map trendradar/."""
from __future__ import annotations

import json
import os
import tempfile
from datetime import timedelta

import pytest

os.environ["TRENDRADAR_DATA"] = tempfile.mkdtemp()

from trendradar import budget, classify, config, db, discovery, snackspert, trends  # noqa: E402
from trendradar.collector import store_posts  # noqa: E402
from trendradar.sources import apify, news  # noqa: E402


REAL_ACCOUNT_USAGE = apify.account_usage


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    """Tests gaan nooit echt naar internet (Apify-tegoed, OpenStreetMap)."""
    from trendradar import geo

    monkeypatch.setattr(apify, "account_usage", lambda: None)

    def no_lookup(place):
        raise geo.requests.ConnectionError("offline in tests")
    monkeypatch.setattr(geo, "lookup", no_lookup)


@pytest.fixture()
def conn(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "t.db")
    db.init()
    c = db.connect()
    yield c
    c.close()


def post(i, text, hours_ago=1, platform="instagram", author=None, **kw):
    return {"id": f"{platform}:{i}", "platform": platform, "author": author or f"maker{i}",
            "text": text, "posted_at": db.iso(db.now() - timedelta(hours=hours_ago)),
            "likes": kw.get("likes", 10), "comments": 1, "shares": 0, "views": 0, **kw}


# --- indelen ---------------------------------------------------------------

def test_relevant_dutch_snack_post():
    c = classify.classify_rules({"text": "Beste frikandel speciaal van Rotterdam, echt lekker! #snackbar #eten010",
                                 "platform": "instagram"})
    assert c["relevance"] > 0.6
    assert c["nl_score"] >= 0.5
    assert "Friet & frituur" in c["categories"]
    assert c["city"] == "Rotterdam"


def test_irrelevant_post_scores_low():
    c = classify.classify_rules({"text": "Great sunset today, love this city #instagood", "platform": "instagram"})
    assert c["relevance"] < 0.2
    assert c["nl_score"] < 0.4


def test_new_supermarket_product():
    c = classify.classify_rules({"text": "Nieuw bij Albert Heijn: chips met kapsalonsmaak, smaaktest!",
                                 "platform": "tiktok"})
    assert "Albert Heijn" in c["brands"]
    assert "Supermarkt & nieuwe producten" in c["categories"]
    assert c["relevance"] >= 0.8


def test_brand_matching_is_strict():
    c = classify.classify_rules({"text": "Dubai chocolade reep #dubaichocolade", "platform": "instagram"})
    assert "Ola" not in c["brands"]
    c = classify.classify_rules({"text": "De nieuwe Lays paprika is top", "platform": "instagram"})
    assert "Lay's" in c["brands"]


def test_compound_words_and_snackbar_not_chips():
    c = classify.classify_rules({"text": "Kipburger bij de snackbar", "platform": "instagram"})
    assert "Burgers" in c["categories"]
    assert "Chips & snacks" not in c["categories"]


def test_phrases_need_a_food_word():
    phrases = classify.key_phrases("Iedereen praat over de pistache kroket. Klinkt raar maar lekker")
    assert "pistache kroket" in phrases
    assert "iedereen praat" not in phrases


def test_city_from_area_code_hashtag():
    c = classify.classify_rules({"text": "Lekker gegeten #eten020", "platform": "instagram"})
    assert c["city"] == "Amsterdam"


def test_merge_ai_overrides_relevance():
    rule = classify.classify_rules({"text": "Frikandel!", "platform": "instagram"})
    merged = classify.merge_ai(rule, {"relevant": False, "relevantie": 0.9, "nederlands": True,
                                      "categorieen": [], "producten": [], "zaak": "", "stad": "",
                                      "sentiment": "neutraal", "samenvatting": ""})
    assert merged["relevance"] < rule["relevance"]
    merged = classify.merge_ai(rule, {"relevant": True, "relevantie": 1, "nederlands": True,
                                      "categorieen": ["Friet & frituur"],
                                      "producten": [{"naam": "Frikandel XXL", "merk": "Mora",
                                                     "winkel": "", "nieuw": True}],
                                      "zaak": "Snackbar Test", "stad": "Utrecht",
                                      "sentiment": "positief", "samenvatting": "x"})
    assert merged["products"][0]["naam"] == "Frikandel XXL"
    assert "Snackbar Test" in merged["places"]
    assert merged["city"] == "Utrecht"


# --- databronnen omzetten ----------------------------------------------------

def test_normalize_instagram_item():
    p = apify.normalize("instagram", {
        "id": "123", "shortCode": "abc", "caption": "Kapsalon! #kapsalon @snackbar_x",
        "ownerUsername": "Foodie", "timestamp": "2026-10-19T10:00:00.000Z", "likesCount": 50,
        "commentsCount": 4, "hashtags": ["kapsalon"], "mentions": ["snackbar_x"],
        "url": "https://www.instagram.com/p/abc/", "locationName": "Snackbar X"})
    assert p["id"] == "instagram:123"
    assert p["author"] == "foodie"
    assert p["posted_at"] == "2026-10-19T10:00:00Z"
    assert p["likes"] == 50 and p["location"] == "Snackbar X"


def test_normalize_tiktok_item():
    p = apify.normalize("tiktok", {
        "id": "999", "text": "Nieuw bij AH", "createTime": 1760868000,
        "authorMeta": {"name": "snacktok", "fans": 1000}, "diggCount": 5, "shareCount": 2,
        "playCount": 100, "commentCount": 1, "hashtags": [{"name": "NieuwBijAH"}],
        "webVideoUrl": "https://tiktok.com/x"})
    assert p["id"] == "tiktok:999"
    assert p["hashtags"] == ["nieuwbijah"]
    assert p["views"] == 100 and p["author_followers"] == 1000


def test_normalize_facebook_item_without_date_is_skipped():
    assert apify.normalize("facebook", {"postId": "1", "text": "x"}) is None
    p = apify.normalize("facebook", {"postId": "1", "text": "Nieuw!", "time": "2026-10-19T10:00:00Z",
                                     "pageName": "AlbertHeijn", "likes": 3})
    assert p["id"] == "facebook:1" and p["author"] == "albertheijn"


def test_parse_google_news_feed():
    xml = """<rss><channel><item><title>Nieuwe chips bij Jumbo - Distrifood</title>
    <link>https://example.com/a</link><pubDate>Mon, 19 Oct 2026 08:00:00 GMT</pubDate>
    <source url="https://distrifood.nl">Distrifood</source></item></channel></rss>"""
    items = news.parse_feed(xml, "nieuw bij Jumbo")
    assert items[0]["outlet"] == "Distrifood"
    assert items[0]["published_at"] == "2026-10-19T08:00:00Z"


# --- budget --------------------------------------------------------------------

def test_budget_plan_spreads_over_remaining_runs(conn):
    from datetime import datetime, timezone

    settings = dict(config.DEFAULTS, maandbudget_eur=46.0, eur_per_usd=0.92)
    start = datetime(2026, 10, 1, 0, 1, tzinfo=timezone.utc)
    p = budget.plan(conn, settings, start)
    assert p["budget_usd"] == 50.0
    total_results = sum(v["results"] for v in p["platforms"].values())
    assert total_results > 0
    assert p["per_run_usd"] * p["runs_left"] == pytest.approx(50.0, rel=0.01)
    # Stond de laptop de hele maand uit, dan krijgt de laatste dag niet het hele restbudget:
    # hooguit twee keer een normale ronde.
    late = budget.plan(conn, settings, datetime(2026, 10, 31, 12, 0, tzinfo=timezone.utc))
    assert late["per_run_usd"] == pytest.approx(2 * p["per_run_usd"], rel=0.02)
    budget.record(conn, "apify", 60.0, "tiktok")
    p = budget.plan(conn, settings)
    assert p["left_usd"] == 0 and all(v["results"] == 0 for v in p["platforms"].values())


def test_allocate_prefers_productive_and_stale_sources():
    now = db.now()
    src = [
        {"id": 1, "yield_score": 0.9, "priority": 1, "last_fetched": db.iso(now - timedelta(hours=30))},
        {"id": 2, "yield_score": 0.1, "priority": 1, "last_fetched": db.iso(now)},
        {"id": 3, "yield_score": None, "priority": 1, "last_fetched": None},
    ]
    chosen = budget.allocate(src, 40, 10, 40, now)
    ids = [s["id"] for s, _ in chosen]
    assert ids[0] in (1, 3) and 2 not in ids[:2]
    assert sum(n for _, n in chosen) <= 40
    assert budget.allocate(src, 5, 10, 40, now) == []


# --- trends & ontdekken ------------------------------------------------------

def test_trend_detects_rising_topic(conn):
    posts = []
    # Basislijn: 'kroket' elke dag een beetje.
    for d in range(1, 7):
        posts.append(post(f"b{d}", "Lekkere kroket bij de snackbar #kroket", hours_ago=24 * d + 3))
    # Vandaag: veel makers over #pistachekroket.
    for i in range(6):
        posts.append(post(f"n{i}", "Pistache kroket bij de snackbar in Rotterdam! #pistachekroket #kroket",
                          hours_ago=2 + i))
    store_posts(conn, posts, None)
    conn.execute("INSERT INTO runs(started_at, status) VALUES (?, 'klaar')",
                 (db.iso(db.now() - timedelta(days=10)),))
    r = trends.compute(conn, "24u")
    names = [t["name"] for t in r["topics"]]
    assert "pistache kroket" in names
    top = next(t for t in r["topics"] if t["name"] == "pistache kroket")
    assert top["new"] is True and top["authors"] == 6
    # Hashtags zijn hulpmiddel om te zoeken, geen onderwerp.
    assert not any(t["name"].startswith("#") for t in r["topics"])
    assert r["history_ok"]


def test_single_author_spam_is_not_a_trend(conn):
    store_posts(conn, [post(f"s{i}", "Koop mijn frikandel #spamfrikandel", author="spammer",
                            hours_ago=1 + i) for i in range(10)], None)
    r = trends.compute(conn, "24u")
    assert "#spamfrikandel" not in [t["name"] for t in r["topics"]]


def test_discovery_adds_hashtag_and_pauses_bad_source(conn):
    store_posts(conn, [post(f"d{i}", "Birria tacos bij de foodtruck, lekker! #birriatacos #foodtruck",
                            hours_ago=3 + i) for i in range(4)], None)
    settings = dict(config.DEFAULTS)
    added = discovery.run(conn, settings)
    values = {(a["kind"], a["value"]) for a in added}
    assert ("hashtag", "birriatacos") in values
    row = conn.execute("SELECT status, origin FROM sources WHERE value = 'birriatacos'").fetchone()
    assert row["status"] == "actief" and row["origin"] == "ontdekt"

    conn.execute("INSERT INTO sources(platform, kind, value, origin, added_at) "
                 "VALUES ('tiktok', 'hashtag', 'saai', 'ontdekt', ?)", (db.iso(db.now()),))
    sid = conn.execute("SELECT id FROM sources WHERE value = 'saai'").fetchone()["id"]
    for _ in range(3):
        discovery.update_source_stats(conn, sid, 30, 1)
    assert conn.execute("SELECT status FROM sources WHERE id = ?", (sid,)).fetchone()["status"] == "gepauzeerd"


def test_empty_account_becomes_invalid_but_manual_stays(conn):
    now = db.iso(db.now())
    conn.execute("INSERT INTO sources(platform, kind, value, origin, added_at) VALUES "
                 "('instagram', 'account', 'bestaatniet', 'start', ?), "
                 "('instagram', 'account', 'mijnfavoriet', 'handmatig', ?)", (now, now))
    for value in ("bestaatniet", "mijnfavoriet"):
        sid = conn.execute("SELECT id FROM sources WHERE value = ?", (value,)).fetchone()["id"]
        discovery.update_source_stats(conn, sid, 0, 0)
        discovery.update_source_stats(conn, sid, 0, 0)
    status = dict(conn.execute("SELECT value, status FROM sources").fetchall())
    assert status["bestaatniet"] == "ongeldig"
    assert status["mijnfavoriet"] == "actief"


def test_snackspert_name_matching():
    reviewed = [{"name": "Snackbar De Gouden Hoek", "url": "u", "norm": snackspert.core_name("Snackbar De Gouden Hoek")}]
    assert snackspert.match("@degoudenhoek_snackbar", reviewed)
    assert snackspert.match("Cafetaria De Gouden Hoek", reviewed)
    assert snackspert.match("Frituur Bakkie Patat", reviewed) is None


def test_demo_round_and_api(conn, monkeypatch):
    from trendradar import demo
    from trendradar.web import app

    n = demo.simulate_round(conn)
    conn.commit()
    assert n > 500
    client = app.test_client()
    for url in ("/api/status", "/api/trends?window=nu", "/api/trends?window=7d&platform=tiktok",
                "/api/posts?q=kroket", "/api/supermarkt", "/api/zaken", "/api/sources",
                "/api/budget", "/api/runs", "/api/export/trends.csv",
                "/api/topic?type=product&name=x", "/api/topic?type=stad&name=Rotterdam"):
        assert client.get(url).status_code == 200, url
    assert client.get("/api/trends?window=1j").status_code == 400
    assert client.post("/api/settings", json={"interval_uren": 0}).status_code == 400
    r = client.post("/api/sources", json={"platform": "tiktok", "kind": "hashtag", "value": "#Kroket"})
    assert r.status_code == 200
    with db.session() as c2:
        assert c2.execute("SELECT 1 FROM sources WHERE value='kroket' AND origin='handmatig'").fetchone()
    demo.clear(conn)
    conn.commit()
    assert conn.execute("SELECT COUNT(*) FROM posts WHERE demo = 1").fetchone()[0] == 0


def test_ai_classification_request_and_parsing(monkeypatch):
    import json as _json
    import types

    import anthropic

    calls = {}

    class FakeMessages:
        def create(self, **kw):
            calls.update(kw)
            body = {"posts": [{"id": "instagram:1", "relevant": True, "relevantie": 0.9,
                               "nederlands": True, "categorieen": ["Chips & snacks"],
                               "producten": [{"naam": "Kapsalon Chips", "merk": "Lay's",
                                              "winkel": "Albert Heijn", "nieuw": True}],
                               "zaak": "", "stad": "", "sentiment": "positief",
                               "samenvatting": "Nieuwe chips getest"}]}
            return types.SimpleNamespace(
                stop_reason="end_turn",
                usage=types.SimpleNamespace(input_tokens=1000, output_tokens=500),
                content=[types.SimpleNamespace(type="text", text=_json.dumps(body))])

    class FakeClient:
        def __init__(self, **kw):
            self.messages = FakeMessages()
            self.beta = types.SimpleNamespace(messages=FakeMessages())

    monkeypatch.setattr(anthropic, "Anthropic", FakeClient)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test")
    settings = dict(config.DEFAULTS)
    results, cost = classify.classify_ai([{"id": "instagram:1", "platform": "instagram",
                                           "text": "Nieuwe chips!"}], settings)
    assert results["instagram:1"]["producten"][0]["naam"] == "Kapsalon Chips"
    assert cost == pytest.approx((1000 * 4 + 500 * 20) / 1_000_000)
    assert calls["model"] == "claude-opus-5-5"
    assert calls["fallbacks"] == "default"
    assert calls["output_config"]["format"]["type"] == "json_schema"
    # Het echte SDK moet deze parameters accepteren.
    import inspect
    sig = inspect.signature(anthropic.resources.beta.messages.Messages.create)
    for name in ("betas", "fallbacks", "output_config", "system", "model", "max_tokens", "messages"):
        assert name in sig.parameters


def test_secrets_are_redacted(monkeypatch):
    monkeypatch.setenv("APIFY_TOKEN", "apify_api_geheim123")
    assert "geheim" not in config.redact("fout bij ...?token=apify_api_geheim123&x=1")


# --- signalen, virale makers, supermarkt, kosten ------------------------------

def _p(i, views, followers, hours_ago, platform="tiktok", **kw):
    return post(f"s{i}", kw.pop("text", "Lekkere frikandel bij de snackbar #frikandel"),
                hours_ago=hours_ago, platform=platform, views=views, author_followers=followers,
                snapshot_at=db.iso(db.now()), **kw)


def test_signals_compare_with_posts_of_same_age(conn):
    from trendradar import signals
    # Oude posts hebben veel views (ze staan er al lang), nieuwe weinig: dat is normaal.
    store_posts(conn, [_p(i, 50_000, 5000, 60) for i in range(20)]
                + [_p(100 + i, 1_000, 5000, 2) for i in range(20)], None)
    base = signals.baselines(conn)
    fresh_normal = signals.of_post(_p(900, 1_000, 5000, 2) | {"fetched_at": db.iso(db.now())}, base)
    assert fresh_normal["views_x"] == 1.0 and not fresh_normal["breakout"]
    viral = signals.of_post(_p(901, 40_000, 900, 2) | {"fetched_at": db.iso(db.now())}, base)
    assert viral["breakout"] and viral["small_creator"] and viral["views_x"] >= 30
    assert "doorbraak" in signals.labels(viral)


def test_breakout_from_unknown_creator_becomes_source_first(conn):
    store_posts(conn, [_p(i, 1_000, 5000, 2 + i % 3) for i in range(20)], None)
    store_posts(conn, [_p(500, 80_000, 700, 3, text="Frikandelijs?! Moet je proberen #frikandelijs #snackbar",
                         author="kleine_maker")], None)
    added = discovery.run(conn, dict(config.DEFAULTS, max_nieuwe_bronnen_per_run=2))
    first = [(a["kind"], a["value"]) for a in added if a["status"] == "actief"]
    assert ("hashtag", "frikandelijs") in first or ("account", "kleine_maker") in first
    assert any("viraal" in a["reason"] for a in added)


def test_trends_show_single_viral_creator(conn):
    store_posts(conn, [_p(i, 1_000, 5000, 1 + i % 4) for i in range(20)], None)
    store_posts(conn, [_p(600, 90_000, 800, 2, text="Bitterbal softijs gespot! #bitterbalijs",
                         author="eenling")], None)
    r = trends.compute(conn, "24u")
    names = [t["name"] for t in r["topics"]]
    assert "bitterbal softijs" in names or "bitterbalijs" in names
    assert r["totals"]["breakouts"] >= 1 and r["breakouts"][0]["signals"]["followers"] == 800


def test_product_names_and_early_signals(conn):
    from trendradar import products
    assert products.same("Kapsalon Chips", "Smaakmakers Kapsalon Chips 150g")
    assert not products.same("Paprika Chips", "Naturel Chips")
    c = classify.classify_rules({"text": "Nieuw bij AH: Paprika Chips Limoen! #ahnieuw", "platform": "tiktok"})
    assert c["products"][0]["naam"] == "Paprika Chips Limoen" and c["products"][0]["winkel"] == "Albert Heijn"
    assert classify.classify_rules({"text": "Nieuw bij AH: mijn collega Jan", "platform": "tiktok"})["products"] == []

    now = db.iso(db.now())
    conn.execute("INSERT INTO products(id, store, name, first_seen, last_seen, is_new) "
                 "VALUES ('ah:1', 'Albert Heijn', 'AH Paprika Chips Limoen', ?, ?, 1)", (now, now))
    store_posts(conn, [post(f"pr{i}", "Nieuw bij AH: Paprika Chips Limoen! Echt lekker", hours_ago=3)
                       for i in range(3)]
                + [post(f"bj{i}", "Gespot bij Jumbo: Bitterbal Chips! Nieuwe smaak", hours_ago=2)
                   for i in range(2)], None)
    o = products.overview(conn)
    assert o["products"][0]["confirmed_by"] == ["winkel", "social"]
    assert o["products"][0]["social_makers"] == 3
    early = {m["name"]: m for m in o["early"]}
    assert "Bitterbal Chips" in early and early["Bitterbal Chips"]["store_tracked"]
    assert "Paprika Chips Limoen" not in early


def test_apify_round_saves_money(conn, monkeypatch):
    """Datumfilter, kosten per bron, en bronnen met weinig nieuws later opnieuw."""
    from trendradar import collector
    from trendradar.sources import apify as ap

    conn.execute("INSERT INTO sources(platform, kind, value, added_at) VALUES "
                 "('tiktok', 'hashtag', 'frikandel', ?), ('tiktok', 'hashtag', 'kroket', ?)",
                 (db.iso(db.now()), db.iso(db.now())))
    conn.commit()
    calls = []

    def fake_run(platform, kind, values, per_value, max_usd, since=None, **kw):
        calls.append({"kind": kind, "values": values, "since": since, "max_usd": max_usd})
        items = [{"id": f"{v}{i}", "text": f"Lekkere {v} bij de snackbar #{v}", "createTime": 1760868000,
                  "authorMeta": {"name": f"m{i}", "fans": 100}, "playCount": 10,
                  "hashtags": [{"name": v}]} for v in values for i in range(per_value)]
        return items, 0.50

    monkeypatch.setattr(ap, "available", lambda: True)
    monkeypatch.setattr(ap, "run_actor", fake_run)
    settings = dict(config.DEFAULTS, maandbudget_eur=600.0)
    run = collector.Run("test")
    collector.step_apify(run, settings, set())
    assert calls and calls[0]["since"]  # datumfilter meegestuurd
    assert all(c["max_usd"] > 0 for c in calls)
    rows = {r["value"]: dict(r) for r in conn.execute("SELECT * FROM sources")}
    assert sum(r["cost_usd"] for r in rows.values()) == pytest.approx(0.5 * len(calls))
    assert all(r["useful_total"] > 0 and r["next_due"] for r in rows.values())

    # Tweede ronde direct erna: bronnen zijn nog niet aan de beurt.
    calls.clear()
    collector.step_apify(run, settings, set())
    assert calls == []

    # Na de wachttijd: zelfde posts terug = 0% nieuw -> volgende keer pas na 4 rondes.
    conn.execute("UPDATE sources SET next_due = ?", (db.iso(db.now() - timedelta(minutes=1)),))
    conn.commit()
    collector.step_apify(run, settings, set())
    r = conn.execute("SELECT next_due, last_new FROM sources WHERE value = 'frikandel'").fetchone()
    assert r["last_new"] == 0
    assert db.parse_iso(r["next_due"]) > db.now() + timedelta(hours=20)


def test_migration_adds_new_columns(tmp_path, monkeypatch):
    import sqlite3
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "oud.db")
    old = sqlite3.connect(tmp_path / "oud.db")
    old.execute("CREATE TABLE sources (id INTEGER PRIMARY KEY, platform TEXT, kind TEXT, value TEXT, "
                "status TEXT, origin TEXT, added_at TEXT)")
    old.commit()
    old.close()
    db.init()
    with db.session() as c:
        cols = {r["name"] for r in c.execute("PRAGMA table_info(sources)")}
    assert {"next_due", "cost_usd", "useful_total", "last_new"} <= cols


def test_first_real_round_after_demo_does_not_lock(conn, monkeypatch):
    """Eerste echte ronde na de demo: demodata wissen en loggen mag niet vastlopen."""
    from trendradar import collector, demo
    from trendradar.sources import apify as ap

    demo.simulate_round(conn)
    conn.commit()
    monkeypatch.setattr(db, "connect", _fast_connect(db.connect))
    monkeypatch.setattr(ap, "available", lambda: True)
    monkeypatch.setattr(ap, "run_actor", lambda *a, **k: ([], 0.0))
    for step in ("step_news", "step_supermarkets", "step_snackspert", "step_remeasure", "step_ai"):
        monkeypatch.setattr(collector, step, lambda *a, **k: None)
    collector.run_once("test")
    with db.session() as c:
        run = collector.last_run(c)
        demo_left = c.execute("SELECT COUNT(*) FROM posts WHERE demo = 1").fetchone()[0]
    assert run["status"] != "fout", run["log"]
    assert demo_left == 0


def _fast_connect(original):
    # Kortere wachttijd, zodat een vastlopende database de test niet 30 s ophoudt.
    def connect():
        c = original()
        c.execute("PRAGMA busy_timeout = 1000")
        return c
    return connect


def test_supermarket_new_products_via_checkjebon(conn, monkeypatch):
    from trendradar import collector
    from trendradar.sources import supermarket

    raw = [{"n": "ah", "u": "https://www.ah.nl/producten/product/", "d": [
               {"n": "AH Naturel chips", "l": "wi1/chips", "p": 1.29, "s": "200 g"}]},
           {"n": "lidl", "u": "https://x/", "d": [{"n": "Kip gyros", "l": "77", "p": 6.19, "s": "750 g"}]},
           {"n": "aldi", "u": "https://y/", "d": []}]
    first = supermarket.parse(raw)
    assert first["Albert Heijn"][0]["url"] == "https://www.ah.nl/producten/product/wi1/chips"
    second = supermarket.parse(raw)
    second["Albert Heijn"] += [
        {"id": "ah:wi2/kapsalon", "store": "Albert Heijn", "name": "AH Kapsalon chips", "size": "150 g",
         "price": 1.99, "url": None},
        {"id": "ah:wi3/shampoo", "store": "Albert Heijn", "name": "AH Shampoo kokos", "size": "300 ml",
         "price": 2.49, "url": None}]
    responses = iter([(first, '"a"'), (None, '"a"'), (second, '"b"')])
    monkeypatch.setattr(supermarket, "fetch", lambda etag=None: next(responses))
    settings = dict(config.DEFAULTS)

    run = collector.Run("test")
    collector.step_supermarkets(run, settings, force=True)   # nulmeting: niets nieuw
    assert conn.execute("SELECT COUNT(*) FROM products WHERE is_new = 1").fetchone()[0] == 0
    collector.step_supermarkets(run, settings, force=True)   # 304: niets veranderd
    assert "geen nieuwe" in run.lines[-1]["tekst"]
    collector.step_supermarkets(run, settings, force=True)
    new = {r["name"]: r["relevant"] for r in conn.execute("SELECT * FROM products WHERE is_new = 1")}
    assert new == {"AH Kapsalon chips": 1, "AH Shampoo kokos": 0}

    from trendradar import products
    o = products.overview(conn)
    assert [p["name"] for p in o["products"]] == ["AH Kapsalon chips"]
    assert len(products.overview(conn, only_relevant=False)["products"]) == 2
    with db.session() as c:
        status = db.get_state(c, "supermarkt_status")
    assert status["Albert Heijn"]["ok"] and not status["Aldi"]["ok"]


def test_apify_minimum_charge_and_cost_estimate(monkeypatch):
    from trendradar.sources import apify as ap

    sent = {}

    class Resp:
        def __init__(self, data, status=200):
            self._d, self.status_code, self.ok = data, status, status < 400
            self.text = ""

        def json(self):
            return self._d

        def raise_for_status(self):
            pass

    def post(url, params=None, headers=None, json=None, timeout=None):
        sent.update(params or {})
        sent["auth"] = headers.get("Authorization")
        return Resp({"data": {"id": "r1", "status": "SUCCEEDED", "defaultDatasetId": "d1"}})

    def get(url, params=None, headers=None, timeout=None):
        if "/datasets/" in url:
            return Resp([{"id": i} for i in range(40)])
        return Resp({"data": {"id": "r1", "status": "SUCCEEDED", "usageTotalUsd": 0}})

    monkeypatch.setenv("APIFY_TOKEN", "tok")
    monkeypatch.setattr(ap.requests, "post", post)
    monkeypatch.setattr(ap.requests, "get", get)
    items, cost = ap.run_actor("tiktok", "hashtag", ["kroket"], 40, 0.05, price_per_1000=1.70)
    assert sent["maxTotalChargeUsd"] == "0.50" and sent["maxItems"] == 40
    assert "token" not in sent and sent["auth"] == "Bearer tok"
    assert cost == pytest.approx(40 * 1.70 / 1000)


def test_only_dutch_posts_count():
    def nl(text):
        return classify.classify_rules({"text": text, "platform": "instagram"})["nl_score"]
    # Spaans, Portugees en Engels met #snackbar: niet Nederlands.
    assert nl("XV años Estefanía Muchas Gracias a Elena por permitirnos ser parte de este día #snackbar") < 0.4
    assert nl("Crocante por fora, irresistível por dentro! As nossas coxinhas são uma tentação #snackbar") < 0.4
    assert nl("A colorful snapshot in the middle of Dirkshorn. The tree now matches the snack bar.") < 0.4
    assert nl("Patatas bravas con salsa de la casa #snackbar #tapas") < 0.4
    # Nederlands, ook kort of gemengd met een Engelse zin.
    assert nl("Honger tijdens Hardenberg Swingt? Sevens has got you covered! Deel dit bericht met je vrienden!") >= 0.4
    assert nl("Bitterbal softijs gespot! #bitterbalijs") >= 0.4
    assert nl("Even iets fris! Niet iedere lunch hoeft een broodje te zijn.") >= 0.4


def test_existing_posts_are_rescored_on_start(conn):
    from trendradar import collector
    store_posts(conn, [post("es1", "Muchas gracias por este día con la familia #snackbar")], None)
    conn.execute("UPDATE posts SET nl_score = 0.9")  # zoals de oude regels het zagen
    db.set_state(conn, "classify_version", 1)
    conn.commit()
    assert collector.reclassify_if_needed() == 1
    assert conn.execute("SELECT nl_score FROM posts").fetchone()[0] < 0.4
    assert collector.reclassify_if_needed() == 0


def test_news_list_only_shows_relevant_news(conn):
    from trendradar import products
    now = db.iso(db.now())
    for nid, title in (("nieuws:a", "Nieuwe chips met kapsalonsmaak bij Albert Heijn"),
                       ("nieuws:b", "Drie doden door stormweer in Frankrijk")):
        conn.execute("INSERT INTO news(id, title, published_at, fetched_at) VALUES (?, ?, ?, ?)",
                     (nid, title, now, now))
        store_posts(conn, [{"id": nid, "platform": "nieuws", "text": title, "posted_at": now}], None)
    titles = [n["title"] for n in products.overview(conn)["news"]]
    assert titles == ["Nieuwe chips met kapsalonsmaak bij Albert Heijn"]


def test_kapsalon_at_the_hairdresser_is_not_a_snack():
    hair = classify.classify_rules({"text": "Nieuwe coupe gezet bij onze kapsalon! #kapper #kapsalon",
                                    "platform": "instagram"})
    snack = classify.classify_rules({"text": "Beste kapsalon van Rotterdam, extra knoflooksaus #kapsalon",
                                     "platform": "instagram"})
    assert hair["relevance"] < 0.35 and snack["relevance"] > 0.7


def test_discovery_only_adds_food_hashtags(conn):
    from trendradar.discovery import acceptable_hashtag
    for bad in ("kapper", "deals", "herfst", "groningen", "albertheijn", "snackbar"):
        assert not acceptable_hashtag(bad), bad
    for good in ("birriatacos", "frikandelijs", "pistachekroket", "foodtruckfestival"):
        assert acceptable_hashtag(good), good
    assert acceptable_hashtag("tanghulu", viral=True)  # nieuwe snack uit een virale post

    store_posts(conn, [post(f"h{i}", "Lekkere frikandel bij de snackbar! #herfst #deals #birriatacos",
                            hours_ago=2 + i) for i in range(4)], None)
    added = {a["value"] for a in discovery.run(conn, dict(config.DEFAULTS)) if a["kind"] == "hashtag"}
    assert "birriatacos" in added and not {"herfst", "deals"} & added


def test_bad_discovered_sources_are_cleaned_up(conn):
    from trendradar import collector
    now = db.iso(db.now())
    for value, reason in (("kapper", "in 5 relevante posts"), ("birriatacos", "in 5 relevante posts"),
                          ("tanghulu", "uit virale post van @x (12× views)")):
        conn.execute("INSERT INTO sources(platform, kind, value, status, origin, reason, added_at) "
                     "VALUES ('instagram', 'hashtag', ?, 'actief', 'ontdekt', ?, ?)", (value, reason, now))
    db.set_state(conn, "classify_version", 2)
    conn.commit()
    collector.reclassify_if_needed()
    status = dict(conn.execute("SELECT value, status FROM sources").fetchall())
    assert status == {"kapper": "afgewezen", "birriatacos": "actief", "tanghulu": "actief"}


def test_indonesian_and_turkish_posts_are_not_dutch():
    def nl(text):
        return classify.classify_rules({"text": text, "platform": "instagram"})["nl_score"]
    assert nl("Enak banget kroket ini, harga murah dan rasa mantap! #snackbar") < 0.4
    assert nl("Çok lezzetli döner, afiyet olsun! #kapsalon") < 0.4
    # Nederlandse post over Indisch eten blijft Nederlands.
    assert nl("Heerlijke nasi en saté gehaald bij de toko, echt een aanrader") >= 0.4


def test_new_compound_snack_words_are_topics(conn):
    assert classify.new_compounds(["frikandelijs", "frikandel", "snackbarlucy", "foodspotsamsterdam",
                                   "kapsalonchips", "lekker"]) == ["frikandelijs", "kapsalonchips"]
    store_posts(conn, [post(f"fi{i}", "Moet je proberen! #frikandelijs", hours_ago=1 + i)
                       for i in range(3)], None)
    r = trends.compute(conn, "24u")
    word = next(t for t in r["topics"] if t["name"] == "frikandelijs")
    assert word["type"] == "woord" and word["type_label"] == "Nieuw woord"


def test_new_word_that_is_part_of_a_product_is_not_repeated(conn):
    store_posts(conn, [post(f"kc{i}", "Nieuw bij AH: Kapsalon Chips! #kapsalonchips", hours_ago=1 + i,
                            products=[{"naam": "Smaakmakers Kapsalon Chips", "merk": "", "winkel": "Albert Heijn",
                                       "nieuw": True}]) for i in range(3)], None)
    names = [t["name"] for t in trends.compute(conn, "24u")["topics"]]
    assert "Smaakmakers Kapsalon Chips" in names and "kapsalonchips" not in names


# --- vóór de hype: nieuw op de kaart, combinaties, fase, Amsterdam, Mijn lijst ---

def test_new_on_the_menu_and_unexpected_combination():
    c = classify.classify_rules({"text": "Nieuw op de kaart: broodje met döner en kaassoufflé! #snackbar",
                                 "platform": "tiktok"})
    assert "nieuw_op_kaart" in c["flags"]
    assert "combinatie:kaassoufflé + döner" in c["flags"]
    gewoon = classify.classify_rules({"text": "Frietje met frikandel bij de snackbar, lekker!",
                                      "platform": "tiktok"})
    assert not any(f.startswith("combinatie:") for f in gewoon["flags"])


def test_amsterdam_neighbourhood_counts_as_amsterdam():
    c = classify.classify_rules({"text": "Weekspecial bij onze zaak in De Pijp: soufflé tosti met kaassoufflé",
                                 "platform": "instagram"})
    assert c["city"] == "Amsterdam"
    assert "nieuw_op_kaart" in c["flags"]


def test_phase_of_the_hype():
    assert trends.phase({"growth_pct": -40, "authors": 3, "news": 0}) == "over"
    assert trends.phase({"growth_pct": 50, "authors": 2, "news": 0}) == "vroeg"
    assert trends.phase({"growth_pct": 50, "authors": 8, "news": 0}) == "groeit"
    assert trends.phase({"growth_pct": 50, "authors": 8, "news": 2}) == "piek"


def test_candidates_and_amsterdam_filter(conn):
    store_posts(conn, [
        post("ams", "Nieuw op de kaart in de Jordaan: tosti met kaassoufflé! #amsterdamfood", platform="tiktok"),
        post("rdam", "Nieuw bij ons: broodje döner met kaassoufflé #eten010", platform="tiktok"),
        post("saai", "Lekker frietje gehaald bij de snackbar", platform="tiktok"),
    ], None)
    ids = [c["id"] for c in trends.candidates(conn, include_demo=False)]
    assert "tiktok:ams" in ids and "tiktok:rdam" in ids and "tiktok:saai" not in ids
    ams = [c["id"] for c in trends.candidates(conn, city="Amsterdam", include_demo=False)]
    assert ams == ["tiktok:ams"]
    assert all(p["city"] == "Amsterdam" for p in trends.load_posts(conn, db.now() - timedelta(days=1),
                                                                   db.now(), city="Amsterdam"))


def test_watchlist_api(conn):
    from trendradar.web import app

    client = app.test_client()
    r = client.post("/api/lijst", json={"kind": "zaak", "ref": "Broodjesbar Stoep", "title": "Broodjesbar Stoep",
                                        "city": "Amsterdam"})
    assert r.status_code == 200
    item = client.get("/api/lijst").get_json()["items"][0]
    assert item["status"] == "te proberen"
    assert client.patch(f"/api/lijst/{item['id']}", json={"status": "gepland", "note": "zaterdag"}).status_code == 200
    assert client.patch(f"/api/lijst/{item['id']}", json={"status": "onzin"}).status_code == 400
    item = client.get("/api/lijst").get_json()["items"][0]
    assert (item["status"], item["note"]) == ("gepland", "zaterdag")
    assert "Broodjesbar Stoep" in client.get("/api/export/lijst.csv").get_data(as_text=True)
    for url in ("/api/vandaag", "/api/kandidaten?city=Amsterdam", "/api/nieuws?soort=amsterdam",
                "/api/trends?city=Amsterdam"):
        assert client.get(url).status_code == 200, url
    assert client.delete(f"/api/lijst/{item['id']}").status_code == 200
    assert client.get("/api/lijst").get_json()["items"] == []


def test_seed_upgrade_adds_new_searches_once(conn):
    from trendradar import seeds

    now = db.iso(db.now())
    seeds.install(conn, now)
    conn.execute("DELETE FROM sources WHERE value = 'nieuw op de kaart snackbar'")
    conn.execute("UPDATE sources SET status = 'gepauzeerd' WHERE value = 'snackbar' AND platform = 'tiktok'")
    db.set_state(conn, "seeds_version", 1)
    assert seeds.upgrade(conn, now) == 1
    assert seeds.upgrade(conn, now) == 0
    row = conn.execute("SELECT status, priority FROM sources WHERE value = 'snackbar' AND platform = 'tiktok'"
                       " AND kind = 'hashtag'").fetchone()
    assert row["status"] == "gepauzeerd" and row["priority"] == 0.5


def test_apify_retries_without_optional_fields(monkeypatch):
    from trendradar.sources import apify as ap

    bodies = []

    class Resp:
        def __init__(self, data, status=200):
            self._d, self.status_code, self.ok, self.text = data, status, status < 400, "invalid input"

        def json(self):
            return self._d

        def raise_for_status(self):
            pass

    def post(url, params=None, headers=None, json=None, timeout=None):
        bodies.append(json)
        if len(bodies) == 1:
            return Resp({}, 400)
        return Resp({"data": {"id": "r1", "status": "SUCCEEDED", "defaultDatasetId": "d1"}})

    def get(url, params=None, headers=None, timeout=None):
        if "/datasets/" in url:
            return Resp([])
        return Resp({"data": {"id": "r1", "status": "SUCCEEDED", "usageTotalUsd": 0.01}})

    monkeypatch.setenv("APIFY_TOKEN", "tok")
    monkeypatch.setattr(ap.requests, "post", post)
    monkeypatch.setattr(ap.requests, "get", get)
    ap.run_actor("tiktok", "zoekterm", ["broodje kaassoufflé"], 10, 0.5)
    assert len(bodies) == 2
    assert ap.OPTIONAL_KEYS & bodies[0].keys() and not ap.OPTIONAL_KEYS & bodies[1].keys()


# --- leren van Mijn lijst --------------------------------------------------------

def test_twist_on_a_known_snack():
    for text in ("Loaded kipcorn gehaald bij de snackbar 🔥", "Mijn eerste #loadedkipcorn bij de cafetaria"):
        assert "twist:loaded kipcorn" in classify.classify_rules({"text": text, "platform": "tiktok"})["flags"]
    gewoon = classify.classify_rules({"text": "Loaded friet bij de snackbar", "platform": "tiktok"})
    assert not any(f.startswith("twist:") for f in gewoon["flags"])


def _save(conn, kind, ref, title, status="te proberen"):
    now = db.iso(db.now())
    conn.execute("INSERT INTO watchlist(kind, ref, title, status, added_at, updated_at) VALUES (?,?,?,?,?,?)",
                 (kind, ref, title, status, now, now))


def test_saved_items_teach_vandaag(conn):
    from trendradar import taste

    # Een gewone kipcorn-post zonder bijzondere signalen: eerst geen kandidaat.
    store_posts(conn, [post("oud", "Loaded kipcorn bij de snackbar, echt lekker", hours_ago=30,
                            platform="tiktok", author="eerste"),
                       post("nieuw", "Vandaag loaded kipcorn gehaald bij de cafetaria #snackbar", hours_ago=2,
                            platform="tiktok", author="tweede")], None)
    before = {c["id"]: c for c in trends.candidates(conn, include_demo=False)}
    assert "Past bij jou" not in " ".join(before.get("tiktok:nieuw", {}).get("reasons", []))

    _save(conn, "post", "tiktok:oud", "Loaded kipcorn")
    prof = taste.profile(conn)
    assert prof["twist:loaded kipcorn"] > 0 and prof["snack:kipcorn"] > 0
    after = {c["id"]: c for c in trends.candidates(conn, include_demo=False)}
    assert "tiktok:oud" not in after  # al bewaard
    assert any(r.startswith("Past bij jou") for r in after["tiktok:nieuw"]["reasons"])
    assert after["tiktok:nieuw"]["score"] > before.get("tiktok:nieuw", {"score": 0})["score"]
    learned = taste.summary(prof)
    assert learned[0]["kenmerk"] == "loaded kipcorn" and learned[0]["soort"] == "twist"
    assert [item["kenmerk"] for item in learned].count("loaded kipcorn") == 1


def test_rejected_items_count_against(conn):
    from trendradar import taste

    _save(conn, "onderwerp", "zin|loaded kipcorn", "loaded kipcorn", status="afgevallen")
    assert taste.profile(conn)["zin:loaded kipcorn"] < 0
    bonus, _, _ = taste.match(taste.profile(conn), {"zin:loaded kipcorn"})
    assert bonus < 0


def test_saved_items_become_sources(conn):
    _save(conn, "onderwerp", "zin|loaded kipcorn", "loaded kipcorn")
    store_posts(conn, [post("km", "Loaded kipcorn bij de snackbar", platform="tiktok", author="kleinemaker")], None)
    _save(conn, "post", "tiktok:km", "Loaded kipcorn")
    found = discovery.list_searches(conn, discovery._known(conn), limit=5)
    values = {(f["platform"], f["kind"], f["value"]) for f in found}
    assert ("tiktok", "zoekterm", "loaded kipcorn") in values
    assert ("tiktok", "account", "kleinemaker") in values
    assert all(f["reason"].startswith("uit Mijn lijst") for f in found)
    # Tips alleen (gewicht 0,5) worden geen bron.
    conn.execute("INSERT INTO spots(name, place, specialty, added_at, updated_at) "
                 "VALUES ('Snackbar de Poppe', 'Coevorden', 'broodje lammert', 'x', 'x')")
    later = discovery.list_searches(conn, discovery._known(conn), limit=10)
    assert not any("poppe" in f["value"] for f in later)


# --- ♥/👎, mijn posts, tips & routes ------------------------------------------------

def _workbook() -> bytes:
    import io as _io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Posts"
    ws.append(["Maand", "Jaar", "Datum", "Type", "Caption (begin)", "Link", "Weergaven", "Bereik", "Likes",
               "Reacties", "Gedeeld", "Opgeslagen", "Volgers", "Format", "Kanaal", "Locatie", "Herkenbaarheid",
               "Indeling", "Volgers telt mee"])
    rows = [
        ("2026-01", "De grote frikandel uit de supermarkt test is hierrrr! Met friet. 1. AH", 400000, "Supermarkttest", "Ja"),
        ("2026-01", "Friet bij Snackbar de Poppe in Coevorden, broodje lammert!", 300000, "Snackbar review", "Ja"),
        ("2026-02", "Friet special in Utrecht", 250000, "Snackbar review", "Ja"),
        ("2026-02", "Friet met mayo bij @frietboer", 220000, "Snackbar review", "Ja"),
        ("2026-02", "Frietje in Den Haag", 210000, "Snackbar review", "Ja"),
        ("2026-03", "Friet in Leiden, krokant", 205000, "Snackbar review", "Ja"),
        ("2026-03", "Pizza in Amsterdam", 100000, "Streetfood review", "Ja"),
        ("2026-03", "ad | Pizza samenwerking #samenwerking", 90000, "Samenwerking", "Nee"),
    ]
    for i, (month, caption, views, fmt, counts) in enumerate(rows):
        ws.append([month, 2026, f"{month}-1{i}", "IG-reel", caption, "open post", views, views // 2, 100, 10,
                   50 + i, 5, 30 + i, fmt, "Snackbar", "Buiten Amsterdam", "Mainstream", "Gecontroleerd", counts])
        ws.cell(ws.max_row, 6).hyperlink = f"https://www.instagram.com/reel/test{i}/"
    # Een afgeleid tabblad met dezelfde post (ingekorte caption): mag geen dubbele post opleveren.
    dup = wb.create_sheet("Samenwerkingen")
    dup.append(["Maand", "Datum", "Caption (begin)", "Weergaven", "Volgers", "Gedeeld", "Link"])
    dup.append(["2026-03", "2026-03-17", "ad | Pizza samenwerking", 90000, 37, 57, "open post"])
    dup.cell(2, 7).hyperlink = "https://www.instagram.com/reel/test7/"
    spots = wb.create_sheet("Alle Spots (buiten A'dam)")
    spots.append(["Naam", "Locatie", "Specialiteit", "Regio"])
    spots.append(["Snackbar de Poppe", "Coevorden", "broodje lammert (hamburger met kaassoufflé)", "Drenthe"])
    spots.append(["Frietboerderij", "ONBEKEND", None, "Utrecht"])
    ams = wb.create_sheet("Amsterdam")
    ams.append(["Naam", "Specialiteit"])
    ams.append(["Döner Father", "croissant döner"])
    routes = wb.create_sheet("Routes")
    routes.append(["#", "Route", "Regio", "Stop 1", "Stop 2", "Stop 3", "KM Retour", "Rijtijd", "Opmerkingen"])
    routes.append([1, "Drenthe kort", "Drenthe", "Snackbar de Poppe (Coevorden) - broodje lammert",
                   "Snackbar City (Hoogeveen) - gehaktbal", "Pizza Pazza (Norg)", "~400 km", "5u", "test"])
    summary = wb.create_sheet("Samenvatting")
    summary.append([None, None])
    summary.append(["TOTAAL", 3])
    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_own_posts_tips_and_routes(conn):
    import io as _io

    from trendradar import ownposts, spots
    from trendradar.web import app

    client = app.test_client()
    r = client.post("/api/import", data={"bestand": (_io.BytesIO(_workbook()), "overzicht.xlsx")},
                    content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    res = r.get_json()
    assert (res["posts"], res["spots"], res["routes"]) == (8, 3, 1)
    bad = client.post("/api/import", data={"bestand": (_io.BytesIO(b"geen excel"), "x.xlsx")},
                      content_type="multipart/form-data")
    assert bad.status_code == 400

    o = ownposts.overview(conn, "2026")
    assert o["kpis"]["posts"] == 8
    assert o["kpis"]["followers"] == sum(30 + i for i in range(7))  # collab telt niet mee
    assert o["supermarket_tests"][0]["product"] == "frikandel"
    assert o["groups"]["format"][0]["name"] == "Supermarkttest"  # gesorteerd op mediaan views
    assert any("Snackbar review" in i for i in o["insights"])
    assert ownposts.lifts(conn)["friet"]["lift"] > 1
    assert "frietboer" in ownposts.posted_handles(conn)

    listing = spots.listing(conn)
    poppe = next(s for s in listing["spots"] if s["name"] == "Snackbar de Poppe")
    assert poppe["posted"]["date"] == "2026-01-11"  # genoemd in een eigen caption
    assert next(s for s in listing["spots"] if s["name"] == "Frietboerderij")["place"] == ""
    assert next(s for s in listing["spots"] if s["name"] == "Döner Father")["place"] == "Amsterdam"
    route = spots.routes(conn)[0]
    assert route["stops"][0] == {"name": "Snackbar de Poppe", "place": "Coevorden", "specialty": "broodje lammert",
                                 "text": "Snackbar de Poppe (Coevorden) - broodje lammert", "status": "nieuw",
                                 "coords": None}
    assert route["maps"].startswith("https://www.google.com/maps/dir/") and "Pizza%20Pazza%20Norg" in route["maps"]

    # Opnieuw importeren: geen dubbele posts of tips, en je status blijft staan.
    assert client.patch(f"/api/tips/{poppe['id']}", json={"status": "gepland", "note": "zaterdag"}).status_code == 200
    client.post("/api/import", data={"bestand": (_io.BytesIO(_workbook()), "overzicht.xlsx")},
                content_type="multipart/form-data")
    assert conn.execute("SELECT COUNT(*) FROM own_posts").fetchone()[0] == 8
    again = conn.execute("SELECT status, note FROM spots WHERE name = 'Snackbar de Poppe'").fetchone()
    assert (again["status"], again["note"]) == ("gepland", "zaterdag")
    assert conn.execute("SELECT COUNT(*) FROM routes").fetchone()[0] == 1
    for url in ("/api/mijnposts", "/api/mijnposts?jaar=2026", "/api/tips", "/api/routes", "/api/export/tips.csv",
                "/api/vandaag", "/api/zaken"):
        assert client.get(url).status_code == 200, url
    assert client.post("/api/tips", json={"name": "Nieuwe Zaak", "place": "Gouda"}).status_code == 200
    assert client.post("/api/tips", json={}).status_code == 400


def test_tip_in_radar_and_in_candidates(conn):
    from trendradar import spots

    conn.execute("INSERT INTO spots(name, place, specialty, added_at, updated_at) "
                 "VALUES ('Snackbar de Poppe', 'Coevorden', 'broodje lammert', 'x', 'x')")
    store_posts(conn, [post("pop", "Nieuw op de kaart bij Snackbar De Poppe: broodje met kroket en pizza",
                            platform="tiktok", author="poppefan")], None)
    hits = spots.radar_hits(conn, include_demo=False)
    assert list(hits.values())[0]["posts"] == 1
    cand = {c["id"]: c for c in trends.candidates(conn, include_demo=False)}["tiktok:pop"]
    assert any(r == "Past bij jou: Snackbar de Poppe" for r in cand["reasons"])


def test_feedback_teaches_and_hides(conn):
    from trendradar import taste
    from trendradar.web import app

    store_posts(conn, [post("a", "Nieuw op de kaart: broodje kroket met kaassoufflé", platform="tiktok", author="fan"),
                       post("b", "Nieuw op de kaart: friet met frikandel en kipcorn speciaal", platform="tiktok",
                            author="fan"),
                       post("c", "Nieuw op de kaart bij ons: kroket met pindasaus", platform="tiktok", author="ander")],
                None)
    conn.commit()
    client = app.test_client()
    assert client.post("/api/feedback", json={"post_id": "tiktok:a", "value": 1}).status_code == 200
    assert client.post("/api/feedback", json={"post_id": "tiktok:b", "value": -1}).status_code == 200
    assert client.post("/api/feedback", json={"post_id": "tiktok:b", "value": 5}).status_code == 400
    assert client.get("/api/feedback").get_json() == {"tiktok:a": 1, "tiktok:b": -1}
    prof = taste.profile(conn)
    assert prof.sources["likes"] == 1 and prof.sources["dislikes"] == 1
    assert "maker:tiktok:fan" not in prof  # ♥ en 👎 voor dezelfde maker heffen elkaar op
    assert prof["snack:kroket"] > 0 and prof["snack:frikandel"] < 0
    ids = [c["id"] for c in trends.candidates(conn, include_demo=False)]
    assert "tiktok:a" not in ids and "tiktok:b" not in ids  # al beoordeeld
    cand = next(c for c in trends.candidates(conn, include_demo=False) if c["id"] == "tiktok:c")
    assert "Past bij jou: kroket" in cand["reasons"]
    client.post("/api/feedback", json={"post_id": "tiktok:a", "value": 0})
    assert client.get("/api/feedback").get_json() == {"tiktok:b": -1}


# --- viral-potentie ---------------------------------------------------------------

def test_viral_potential_signals():
    from trendradar import potential

    ctx = {"radar": {}, "snacks": {}, "locations": {}, "channels": {}, "recognition": {}, "user": {}}

    def tip(name, specialty, **kw):
        return {"id": 1, "name": name, "place": kw.get("place", "Zwolle"), "specialty": specialty, "lists": [],
                "note": "", "status": "te proberen"}

    koet = potential.score_tip(tip("Snackbar de Poppe", "broodje lammert (hamburger met kaassoufflé)"), ctx)
    kinds = {r["kind"] for r in koet["reasons"]}
    assert {"creatie", "combinatie", "herkenbaar"} <= kinds
    assert "Eigen creatie: broodje lammert" in [r["text"] for r in koet["reasons"]]
    plain = potential.score_tip(tip("Burgerbar", "lekker vlees, goed broodje, geen gekke dingen"), ctx)
    assert plain["score"] < koet["score"] and plain["potential"] < 30
    tank = potential.score_tip(tip("Cafetaria", "Friet Tank (friet in de vorm van een pantservoertuig)"), ctx)
    assert "spektakel" in {r["kind"] for r in tank["reasons"]}
    tipped = potential.score_tip(tip("Burgerbar", "smashburger; twee keer getipt"), ctx)
    assert "getipt" in {r["kind"] for r in tipped["reasons"]}
    # De naam van de zaak is geen snack: 'Cafetaria de Mariënburgher' levert geen 'creatie' op.
    named = potential.score_tip(tip("Cafetaria de Mariënburgher", "snack de Bamia (zelf ontwikkeld)"), ctx)
    assert any(r["kind"] == "creatie" and "bamia" in r["text"] for r in named["reasons"])
    assert not any(r["kind"] == "creatie" and "mariënburgher" in r["text"] for r in named["reasons"])
    # Jouw publiek telt mee.
    ctx2 = {**ctx, "snacks": {"friet": {"lift": 1.3, "posts": 30}}}
    assert potential.score_tip(tip("X", "friet speciaal"), ctx2)["score"] > potential.score_tip(tip("X", "friet speciaal"), ctx)["score"]


def test_viral_potential_ranking_skips_done_and_adds_route_stops(conn):
    from trendradar import potential
    from trendradar.web import app

    now = db.iso(db.now())
    for name, spec, status in (("Snackbar de Poppe", "broodje lammert (hamburger met kaassoufflé)", "te proberen"),
                               ("Al Gedaan", "broodje viagraa met döner", "gereviewd"),
                               ("Gewoon Burgers", "burgers", "te proberen")):
        conn.execute("INSERT INTO spots(name, place, specialty, status, added_at, updated_at) VALUES (?,?,?,?,?,?)",
                     (name, "Coevorden", spec, status, now, now))
    conn.execute("INSERT INTO routes(name, stops, added_at) VALUES (?, ?, ?)",
                 ("Drenthe", json.dumps(["Snackbar de Poppe (Coevorden) - broodje lammert",
                                         "Snackbar de Luifel (Barger-Compascuum) - luifel snack"]), now))
    conn.commit()
    r = potential.ranking(conn, 10, include_demo=False)
    names = [t["name"] for t in r["top"]]
    assert names[0] == "Snackbar de Poppe" and "Al Gedaan" not in names
    luifel = next(t for t in r["top"] if t["name"] == "Snackbar de Luifel")
    assert luifel["id"] is None and luifel["lists"] == ["Uit routes"]
    assert r["skipped"]["status"] == 1
    assert r["routes"][0]["name"] == "Drenthe" and r["routes"][0]["maps"].startswith("https://www.google.com/maps/dir/")
    assert app.test_client().get("/api/potentie?limit=5").status_code == 200


# --- uit de ronde van 23 oktober ---------------------------------------------------

def test_relisted_supermarket_product_is_not_new(conn, monkeypatch):
    from trendradar import collector
    from trendradar.sources import supermarket

    base = [{"id": "spar:chips-1", "store": "Spar", "name": "Lay's Paprika chips", "size": "225 Gram",
             "price": 2.5, "url": None}]
    later = base + [{"id": "spar:chips-2", "store": "Spar", "name": "Lay's paprika chips", "size": "225 gram",
                     "price": 2.5, "url": None},
                    {"id": "spar:kroket-9", "store": "Spar", "name": "Kwekkeboom kroket", "size": "4 stuks",
                     "price": 3.5, "url": None}]
    responses = iter([({"Spar": base}, '"a"'), ({"Spar": later}, '"b"')])
    monkeypatch.setattr(supermarket, "fetch", lambda etag=None: next(responses))
    run = collector.Run("test")
    collector.step_supermarkets(run, dict(config.DEFAULTS), force=True)
    collector.step_supermarkets(run, dict(config.DEFAULTS), force=True)
    new = [r["name"] for r in conn.execute("SELECT name FROM products WHERE is_new = 1")]
    assert new == ["Kwekkeboom kroket"]
    assert "Spar 1 (1)" in run.lines[-1]["tekst"]
    # Al eerder ten onrechte als nieuw gemeld: wordt rechtgezet.
    conn.execute("UPDATE products SET is_new = 1 WHERE id = 'spar:chips-2'")
    assert collector.unmark_relisted(conn) == 1


def test_apify_retries_only_when_safe(monkeypatch):
    import requests as rq

    from trendradar.sources import apify as ap

    monkeypatch.setattr(ap, "RETRY_WAIT_S", 0)
    calls = []

    def flaky(exc):
        def f(*a, **k):
            calls.append(1)
            if len(calls) == 1:
                raise exc
            return "ok"
        return f

    from urllib3.exceptions import MaxRetryError, NewConnectionError, ProtocolError

    dns = rq.ConnectionError(MaxRetryError(None, "/runs", NewConnectionError(None, "Failed to resolve")))
    monkeypatch.setattr(ap.requests, "post", flaky(dns))
    assert ap._start("https://x") == "ok" and len(calls) == 2
    calls.clear()
    # Verbinding viel weg terwijl Apify het verzoek misschien al had: niet nog eens starten.
    dropped = rq.ConnectionError(ProtocolError("Connection aborted.", ConnectionResetError()))
    monkeypatch.setattr(ap.requests, "post", flaky(dropped))
    with pytest.raises(rq.ConnectionError):
        ap._start("https://x")
    assert len(calls) == 1
    calls.clear()
    monkeypatch.setattr(ap.requests, "post", flaky(rq.ReadTimeout("run loopt misschien al")))
    with pytest.raises(rq.ReadTimeout):
        ap._start("https://x")  # niet opnieuw starten: dat zou dubbel kosten
    assert len(calls) == 1
    calls.clear()
    monkeypatch.setattr(ap.requests, "get", flaky(rq.Timeout("traag")))
    assert ap._get("https://x") == "ok" and len(calls) == 2


def test_network_error_does_not_punish_sources(conn, monkeypatch):
    import requests as rq

    from trendradar import collector
    from trendradar.sources import apify as ap

    conn.execute("INSERT INTO sources(platform, kind, value, status, origin, added_at) "
                 "VALUES ('instagram', 'account', 'snackbardepoppe', 'actief', 'ontdekt', 'x')")
    conn.commit()
    monkeypatch.setattr(ap, "available", lambda: True)

    def down(*a, **k):
        raise rq.ConnectionError("Failed to resolve 'api.apify.com'")
    monkeypatch.setattr(ap, "run_actor", down)
    run = collector.Run("test")
    for _ in range(3):
        collector.step_apify(run, dict(config.DEFAULTS), set())
    s = conn.execute("SELECT status, empty_streak FROM sources WHERE value = 'snackbardepoppe'").fetchone()
    assert (s["status"], s["empty_streak"]) == ("actief", 0)
    assert any("geen verbinding met Apify" in line["tekst"] for line in run.lines)


def test_tips_alone_never_become_sources(conn):
    now = db.iso(db.now())
    # Dezelfde zaak in twee lijsten (met een iets andere plaatsnaam) telde eerder dubbel.
    for place in ("Colmschate", "Colmschate (Deventer)"):
        conn.execute("INSERT INTO spots(name, place, specialty, added_at, updated_at) VALUES (?, ?, '', ?, ?)",
                     ("Paul van Gurp", place, now, now))
    assert not discovery.list_searches(conn, discovery._known(conn), limit=10)
    # Een bron die eerder zo is toegevoegd, gaat terug naar kandidaat.
    conn.execute("INSERT INTO sources(platform, kind, value, status, origin, reason, added_at) VALUES "
                 "('tiktok', 'zoekterm', 'paul van gurp', 'actief', 'ontdekt', 'uit Mijn lijst of ♥: Paul van Gurp', ?)",
                 (now,))
    assert discovery.prune_list_sources(conn) == ["paul van gurp"]
    assert conn.execute("SELECT status FROM sources WHERE value = 'paul van gurp'").fetchone()[0] == "kandidaat"


def test_viral_post_hashtags_must_be_about_food():
    assert not discovery.acceptable_hashtag("gerbrandytoren", viral=True)
    assert discovery.acceptable_hashtag("frikandelijs", viral=True)
    assert discovery.acceptable_hashtag("birriatacos")


# --- Apify-tegoed, routes bedenken, nieuwe tips uit de radar ------------------------

def test_apify_account_limit_stops_the_round(conn, monkeypatch):
    from trendradar import collector
    from trendradar.sources import apify as ap

    class Resp:
        ok, status_code = True, 200

        def json(self):
            return {"data": {"current": {"monthlyUsageUsd": 5.02}, "limits": {"maxMonthlyUsageUsd": 5},
                             "monthlyUsageCycle": {"endAt": "2026-11-01T00:00:00.000Z"}}}

    monkeypatch.setenv("APIFY_TOKEN", "tok")
    monkeypatch.setattr(ap.requests, "get", lambda *a, **k: Resp())
    monkeypatch.setattr(ap, "account_usage", REAL_ACCOUNT_USAGE)
    assert ap.account_usage() == {"used_usd": 5.02, "limit_usd": 5.0, "cycle_start": None,
                                  "cycle_end": "2026-11-01T00:00:00.000Z"}
    started = []
    monkeypatch.setattr(ap, "run_actor", lambda *a, **k: started.append(1) or ([], 0.0))
    conn.execute("INSERT INTO sources(platform, kind, value, status, origin, added_at) "
                 "VALUES ('tiktok', 'zoekterm', 'broodje kaassoufflé', 'actief', 'start', 'x')")
    conn.commit()
    run = collector.Run("test")
    collector.step_apify(run, dict(config.DEFAULTS), set())
    assert not started
    assert any("Apify-tegoed op" in line["tekst"] for line in run.lines)
    assert db.get_state(conn, "apify_account")["limit_usd"] == 5.0
    assert budget.overview(conn, dict(config.DEFAULTS))["apify_account"]["used_usd"] == 5.02
    assert ap._is_limit_error(402, "") and ap._is_limit_error(403, "Monthly usage hard limit exceeded")
    assert not ap._is_limit_error(400, "invalid input")


def test_route_proposals_learn_and_can_be_rejected(conn):
    from trendradar import geo, routeplan

    assert geo.clean_place("Colmschate (Deventer)") == "Colmschate, Deventer"
    assert geo.clean_place("ONBEKEND") == ""
    now = db.iso(db.now())
    tips = [("Snackbar de Poppe", "Coevorden", "broodje lammert (hamburger met kaassoufflé)", (52.66, 6.74)),
            ("Snackbar de Luifel", "Barger-Compascuum", "luifel snack, zelfgemaakt", (52.75, 7.04)),
            ("Lunchroom Old Amsterdam", "Nieuw-Amsterdam", "broodje smikkelvlees met joppiesaus", (52.71, 6.86)),
            ("Ver Weg", "Maastricht", "kapsalon met kaassoufflé", (50.85, 5.69))]
    for name, place, spec, (lat, lon) in tips:
        conn.execute("INSERT INTO spots(name, place, specialty, added_at, updated_at) VALUES (?,?,?,?,?)",
                     (name, place, spec, now, now))
        conn.execute("INSERT INTO geo(place, lat, lon, looked_up) VALUES (?,?,?,?)", (place, lat, lon, now))
    conn.execute("INSERT INTO routes(name, km, stops, added_at) VALUES ('kort', '~450 km', ?, ?)",
                 (json.dumps(["A (Nergenshuizen)", "B (Nergenshuizen)", "C (Nergenshuizen)"]), now))
    conn.commit()
    p = routeplan.proposals(conn, include_demo=False)
    assert p["habits"]["stops"] == [3] and p["habits"]["max_km"] == 450
    assert p["missing"] == 1  # alleen Nergenshuizen uit de route is nog onbekend
    routes = p["routes"]
    assert routes, p
    names = {s["name"] for s in routes[0]["stops"]}
    assert names == {"Snackbar de Poppe", "Snackbar de Luifel", "Lunchroom Old Amsterdam"}
    assert "Ver Weg" not in names  # ligt niet bij de andere drie
    assert routes[0]["maps"].startswith("https://www.google.com/maps/dir/")
    # Bewaren zet hem bij je routes, afwijzen laat hem niet meer terugkomen.
    routeplan.save(conn, routes[0])
    assert conn.execute("SELECT COUNT(*) FROM routes WHERE origin = 'bedacht'").fetchone()[0] == 1
    assert not routeplan.proposals(conn, include_demo=False)["routes"]  # bestaat nu al als route
    conn.execute("DELETE FROM routes WHERE origin = 'bedacht'")
    routeplan.reject(conn, routes[0]["key"])
    assert not routeplan.proposals(conn, include_demo=False)["routes"]


def test_geocoding_fills_missing_places(conn, monkeypatch):
    from trendradar import geo

    now = db.iso(db.now())
    conn.execute("INSERT INTO spots(name, place, specialty, added_at, updated_at) VALUES ('X', 'Barger-Compascuum', '', ?, ?)",
                 (now, now))
    conn.commit()
    monkeypatch.setattr(geo, "PAUSE_S", 0)
    monkeypatch.setattr(geo, "lookup", lambda place: (52.75, 7.04))
    assert geo.missing(conn) == ["Barger-Compascuum"]  # Amsterdam (thuisbasis) kent de tool al
    assert geo.fill(limit=10)["found"] == 1
    assert geo.coords(conn, "Barger-Compascuum") == (52.75, 7.04)
    assert geo.missing(conn) == []


def test_radar_suggestions_are_new_places(conn):
    from trendradar import spots

    store_posts(conn, [post(f"s{i}", "Nieuw op de kaart bij @frituurdekoe in Zwolle: kroket met kaassoufflé",
                            hours_ago=2 + i, author=f"maker{i}", platform="tiktok") for i in range(3)], None)
    sugg = spots.radar_suggestions(conn, include_demo=False)
    koe = next(s for s in sugg if s["name"] == "frituurdekoe")
    assert any("nieuw op de kaart" in r for r in koe["reasons"])
    now = db.iso(db.now())
    conn.execute("INSERT INTO spots(name, place, added_at, updated_at) VALUES ('Frituur de Koe', 'Zwolle', ?, ?)",
                 (now, now))
    assert not any(s["name"] == "frituurdekoe" for s in spots.radar_suggestions(conn, include_demo=False))


# --- review: herkenning, import en analyse --------------------------------------------

@pytest.mark.parametrize("text,expected", [
    ("Lekker gegeten #foodamsterdam", "Amsterdam"),        # niet Rotterdam door 'rdam'
    ("Oliebollen! #oliebollen2026", None),                 # 026 is hier een jaartal, niet Arnhem
    ("Friet in een houten bakje", None),
    ("Broodje met jonge gouda", None),
    ("Dit kan leiden tot een snackverslaving", None),
    ("Softijs #ijshoorn", None),
    ("Kroket gehaald in Zuidoost-Brabant", None),
    ("Nieuw in Gouda: kaassoufflé", "Gouda"),
    ("Friet #eten010", "Rotterdam"),
    ("Soufflé tosti in De Pijp", "Amsterdam"),
])
def test_city_detection(text, expected):
    assert classify.find_city(" " + text.lower() + " ", classify.extract_hashtags(text), None) == expected


def test_location_that_is_only_a_place_is_no_business():
    assert not classify.location_is_business("Amsterdam, Netherlands")
    assert not classify.location_is_business("Kalverstraat 12, 1012 NX Amsterdam, Nederland")
    assert not classify.location_is_business("Nederland")
    assert classify.location_is_business("Snackbar De Poppe")
    c = classify.classify_rules({"text": "Nieuwe sneakers en een frietje", "platform": "instagram",
                                 "location": "Amsterdam, Netherlands"})
    assert c["places"] == [] and c["city"] == "Amsterdam"


@pytest.mark.parametrize("text,expected", [
    ("Pita met döner, kaassoufflé en knoflooksaus", "kaassoufflé + döner"),
    ("Kaassoufflé en döner, wat een combi", "kaassoufflé + döner"),
    ("Suikerbroodburger in Sneek!", "burger + suikerbrood"),
    ("croissant döner", "döner + croissant"),
    ("Patatje met, een frikandel en een kipcorn", None),   # een bestelling, geen gerecht
    ("Friet met een milkshake", None),
    ("Kipcorn gehaald @kebabhouse_ankara", None),          # snack in een accountnaam
    ("Broodje kroket bij Burgerbar Smash", None),          # snack in de naam van de zaak
    ("Wij zoeken iemand die frikandellen bakken en een pizza maken kan", None),
])
def test_combination_needs_one_dish(text, expected):
    assert classify.find_combination(" " + text.lower() + " ", classify.extract_hashtags(text)) == expected


def test_language_edge_cases():
    def nl(text):
        return classify.classify_rules({"text": text, "platform": "instagram"})["nl_score"]
    assert nl("Tried the famous Dutch bitterballen in Amsterdam, so good!") < 0.4
    assert nl("Kapsalon met extra kaas, wie durft? Wie wil?") >= 0.4   # 'wie' is ook Nederlands
    assert nl("Frikandellen 🔥🔥") >= 0.4                              # meervoud telt
    assert classify.new_compounds(["frikandellen", "burgerlover", "burgerking", "frikandelijs"]) == ["frikandelijs"]
    assert "twist:mini bitterbal" not in classify.find_flags(" mini bitterballen bij de snackbar ", [], 1.0)


def test_creation_versus_variant():
    from trendradar import potential

    assert potential._creation("broodje lammert (hamburger met kaassoufflé)") == ("creatie", "broodje lammert")
    assert potential._creation("Steenwijkerbal") == ("creatie", "steenwijkerbal")
    assert potential._creation("rundvleeskroket") == ("variant", "rundvleeskroket")
    assert potential._creation("boerenfrikandel") == ("variant", "boerenfrikandel")
    for ordinary in ("broodje carpaccio", "broodje haring", "lekkere snack", "warme snack", "filet americain"):
        assert potential._creation(ordinary) is None, ordinary
    ctx = {"radar": {}, "snacks": {}, "locations": {}, "channels": {}, "recognition": {}, "user": {}}
    tip = {"id": 1, "name": "X", "place": "Zwolle", "lists": [], "status": "te proberen", "note": "",
           "specialty": "geen gekke dingen"}
    kinds = {r["kind"] for r in potential.score_tip(tip, ctx)["reasons"]}
    assert "spektakel" not in kinds and "algemeen" in kinds
    closed = potential.score_tip({**tip, "specialty": "burgers", "note": "tijdelijk gesloten"}, ctx)
    assert "gesloten" in {r["kind"] for r in closed["reasons"]} and closed["potential"] == 0


def test_import_is_robust(conn):
    import io as _io

    import openpyxl

    from trendradar import imports, ownposts

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Export"
    ws.append(["Overzicht Instagram", "bijgewerkt", "2026-10-01"])          # titelregel
    ws.append(["Date", "Caption", "Link", "Views", "Followers"])           # Engelse kolommen
    ws.append(["17-03-2026", "Sommige snacks zijn gewoon top", "https://x/1", "12,5K", "1.234"])
    ws.append(["2026-04-02", "Broodje kroket", "https://x/2", "3.4K", 10])
    ws.append(["", "Totaal", "", "15900", ""])                             # totaalregel
    tips = wb.create_sheet("Buiten Amsterdam")
    tips.append(["#", "Naam", "Plaats", "Specialiteit"])                    # naam niet in de eerste kolom
    tips.append([1, "Paul van Gurp", "Colmschate", "burger"])
    tips.append([2, "Paul van Gurp", "Colmschate (Deventer)", "Colmschate burger, prijswinnaar"])
    routes = wb.create_sheet("Routes")
    routes.append(["Route", "Stop 1", "Stop 2"])
    routes.append(["Oost", "Paul van Gurp (Colmschate) - burger", "Zaak B (Deventer)"])
    buf = _io.BytesIO()
    wb.save(buf)
    parsed = imports.read_workbook(buf.getvalue())
    assert [(p["month"], p["views"], p["followers"]) for p in parsed["posts"]] == [("2026-03", 12500, 1234),
                                                                                   ("2026-04", 3400, 10)]
    assert all(s["place"] != "Amsterdam" for s in parsed["spots"])
    imports.save(conn, parsed)
    spots = conn.execute("SELECT name, specialty FROM spots").fetchall()
    assert len(spots) == 1 and "prijswinnaar" in spots[0]["specialty"]     # dubbele tip samengevoegd
    conn.execute("UPDATE routes SET status = 'bezocht'")
    imports.save(conn, parsed)                                              # opnieuw importeren
    assert conn.execute("SELECT status FROM routes").fetchone()[0] == "bezocht"
    assert ownposts.overview(conn)["kpis"]["posts"] == 2
    # Een oude import met een maand als 'maart 2026' laat de pagina niet crashen.
    conn.execute("UPDATE own_posts SET month = 'maart 2026'")
    assert ownposts.overview(conn)["kpis"]["posts"] == 2


def test_insights_are_escaped(conn):
    from trendradar import ownposts

    conn.execute("INSERT INTO own_posts(id, month, caption, views, format, followers_count, imported_at) "
                 "VALUES ('a', '2026-03', 'De grote <img src=x onerror=alert(1)> uit de supermarkt test', 100, "
                 "'Supermarkttest', 1, 'x')")
    insights = " ".join(ownposts.overview(conn)["insights"])
    assert "<img" not in insights and "&lt;img" in insights


def test_route_save_once_and_nested_place(conn):
    from trendradar import imports, routeplan

    proposal = {"name": "Test", "km": 100, "duration": "1u", "stops": [
        {"name": "Paul van Gurp", "place": "Colmschate (Deventer)", "specialty": "burger"},
        {"name": "Zaak B", "place": "Deventer", "specialty": ""}]}
    first = routeplan.save(conn, proposal)
    assert routeplan.save(conn, proposal) == first                          # dubbelklik: geen tweede route
    stops = json.loads(conn.execute("SELECT stops FROM routes").fetchone()[0])
    assert imports.parse_stop(stops[0]) == {"name": "Paul van Gurp", "place": "Colmschate, Deventer",
                                            "specialty": "burger"}


def test_product_guess_needs_a_store_or_brand():
    from trendradar import products

    assert products.guess_from_text("de nieuwe Smashburger van de week geprobeerd bij Snackbar X",
                                    classify._is_food_word) == []
    found = products.guess_from_text("De nieuwe Doritos Loaded Cheese bij de Jumbo", classify._is_food_word)
    assert found and found[0]["naam"] == "Doritos Loaded Cheese" and found[0]["winkel"] == "Jumbo"


# --- review: geld en ophaalrondes --------------------------------------------------------

def _apify_offline(monkeypatch):
    from trendradar.sources import apify as ap

    monkeypatch.setattr(ap, "available", lambda: True)
    monkeypatch.setattr(ap, "normalize", lambda platform, item: item)
    monkeypatch.setattr(ap, "account_usage", lambda: None)
    return ap


def test_interrupted_run_is_counted_and_recovered(conn, monkeypatch):
    from trendradar import collector

    ap = _apify_offline(monkeypatch)
    conn.execute("INSERT INTO sources(platform, kind, value, status, origin, added_at) "
                 "VALUES ('tiktok', 'zoekterm', 'broodje kaassoufflé', 'actief', 'start', 'x')")
    conn.commit()

    def broken(*a, **k):
        raise ap.ApifyRunFailed("verbinding weg tijdens de run", "run-1", 0.50, "ds-1")
    monkeypatch.setattr(ap, "run_actor", broken)
    run = collector.Run("test")
    collector.step_apify(run, dict(config.DEFAULTS), set())
    assert budget.spent_usd(conn, "apify") == pytest.approx(0.50)          # voorlopig geteld
    assert db.get_state(conn, "apify_pending")[0]["run_id"] == "run-1"
    # Volgende ronde: de run is klaar, kostte $0.12 en leverde een post op.
    monkeypatch.setattr(ap, "run_info", lambda rid: {"status": "SUCCEEDED", "usage_usd": 0.12, "dataset": "ds-1"})
    monkeypatch.setattr(ap, "dataset_items", lambda ds, limit=1000: [
        post("rec1", "Broodje kaassoufflé, eindelijk geprobeerd!", platform="tiktok")])
    monkeypatch.setattr(ap, "run_actor", lambda *a, **k: ([], 0.0))
    collector.step_reconcile_apify(collector.Run("test"), dict(config.DEFAULTS))
    assert budget.spent_usd(conn, "apify") == pytest.approx(0.12)
    assert conn.execute("SELECT COUNT(*) FROM posts WHERE id = 'tiktok:rec1'").fetchone()[0] == 1
    assert db.get_state(conn, "apify_pending") == []


def test_spend_is_kept_when_storing_fails_and_unmatched_posts_are_stored(conn, monkeypatch):
    from trendradar import collector

    ap = _apify_offline(monkeypatch)
    for tag in ("kroket", "frikandel"):
        conn.execute("INSERT INTO sources(platform, kind, value, status, origin, added_at) "
                     "VALUES ('instagram', 'hashtag', ?, 'actief', 'start', 'x')", (tag,))
    conn.commit()
    items = [post(f"u{i}", "Lekker gegeten bij de snackbar", hashtags=[]) for i in range(4)]
    monkeypatch.setattr(ap, "run_actor", lambda *a, **k: (items, 0.40))
    collector.step_apify(collector.Run("test"), dict(config.DEFAULTS), set())
    assert budget.spent_usd(conn, "apify") == pytest.approx(0.40)
    # Posts zonder herkenbare hashtag zijn toch opgeslagen, verdeeld over beide bronnen.
    assert conn.execute("SELECT COUNT(*) FROM posts").fetchone()[0] == 4
    assert {r[0] for r in conn.execute("SELECT DISTINCT source_id FROM posts")} == {1, 2}

    def boom(*a, **k):
        raise RuntimeError("kapotte post")
    monkeypatch.setattr(collector, "store_posts", boom)
    conn.execute("UPDATE sources SET next_due = NULL")
    conn.commit()
    collector.step_apify(collector.Run("test"), dict(config.DEFAULTS), set())
    assert budget.spent_usd(conn, "apify") == pytest.approx(0.80)            # betaald is betaald


def test_offline_round_is_retried_soon(conn, monkeypatch):
    import requests as rq

    from trendradar import collector

    ap = _apify_offline(monkeypatch)
    conn.execute("INSERT INTO sources(platform, kind, value, status, origin, added_at) "
                 "VALUES ('tiktok', 'zoekterm', 'kapsalon', 'actief', 'start', 'x')")
    conn.commit()

    def down(*a, **k):
        raise rq.ConnectionError("geen internet")
    monkeypatch.setattr(ap, "run_actor", down)
    run = collector.Run("test")
    collector.step_apify(run, dict(config.DEFAULTS), set())
    run.finish()
    assert conn.execute("SELECT status FROM runs WHERE id = ?", (run.id,)).fetchone()[0] == "fout"
    settings = dict(config.DEFAULTS)
    assert not collector.due(conn, settings)
    conn.execute("UPDATE runs SET started_at = ?", (db.iso(db.now() - timedelta(minutes=31)),))
    assert collector.due(conn, settings)


def test_quiet_account_is_not_marked_invalid(conn):
    conn.execute("INSERT INTO sources(platform, kind, value, status, origin, added_at, useful_total) "
                 "VALUES ('instagram', 'account', 'snackbardepoppe', 'actief', 'ontdekt', 'x', 3)")
    sid = conn.execute("SELECT id FROM sources").fetchone()[0]
    for _ in range(5):
        discovery.update_source_stats(conn, sid, 0, 0, 0, 0, 0.0, 6, filtered=True)
    assert conn.execute("SELECT status FROM sources WHERE id = ?", (sid,)).fetchone()[0] == "actief"
    for _ in range(2):
        discovery.update_source_stats(conn, sid, 0, 0, 0, 0, 0.0, 6, filtered=False)
    assert conn.execute("SELECT status FROM sources WHERE id = ?", (sid,)).fetchone()[0] == "ongeldig"


def test_parked_candidate_is_promoted_when_it_goes_viral(conn, monkeypatch):
    now = db.iso(db.now())
    conn.execute("INSERT INTO sources(platform, kind, value, status, origin, added_at) "
                 "VALUES ('tiktok', 'account', 'kleinemaker', 'kandidaat', 'ontdekt', ?)", (now,))
    monkeypatch.setattr(discovery, "candidates", lambda c, **k: [
        {"platform": "tiktok", "kind": "account", "value": "kleinemaker", "score": 40,
         "reason": "kleine maker ging viraal"}])
    added = discovery.run(conn, dict(config.DEFAULTS))
    assert added and conn.execute("SELECT status FROM sources WHERE value = 'kleinemaker'").fetchone()[0] == "actief"


def test_account_case_does_not_create_duplicates(conn):
    from trendradar import seeds

    now = db.iso(db.now())
    conn.execute("INSERT INTO sources(platform, kind, value, status, origin, reason, added_at) "
                 "VALUES ('facebook', 'account', 'AlbertHeijn', 'ongeldig', 'start', "
                 "'twee rondes zonder resultaat (bestaat niet of is privé?)', ?)", (now,))
    conn.execute("INSERT INTO sources(platform, kind, value, status, origin, added_at) "
                 "VALUES ('facebook', 'account', 'albertheijn', 'actief', 'ontdekt', ?)", (now,))
    assert ("facebook", "account", "albertheijn") in discovery._known(conn)
    db.set_state(conn, "seeds_version", 2)
    seeds.upgrade(conn, now)
    rows = conn.execute("SELECT value, status FROM sources WHERE platform = 'facebook' "
                        "AND lower(value) = 'albertheijn'").fetchall()
    assert [(r["value"], r["status"]) for r in rows] == [("AlbertHeijn", "actief")]


def test_remeasure_stops_when_budget_is_used(conn, monkeypatch):
    from trendradar import collector

    ap = _apify_offline(monkeypatch)
    called = []
    monkeypatch.setattr(ap, "run_actor", lambda *a, **k: called.append(1) or ([], 0.0))
    store_posts(conn, [post("fast", "Kroket die viraal gaat", platform="tiktok", views=90000, url="https://t/1",
                            hours_ago=10)], None)
    budget.record(conn, "apify", 500.0, "tiktok")
    conn.execute("UPDATE posts SET fetched_at = ?", (db.iso(db.now() - timedelta(hours=8)),))
    conn.commit()
    collector.step_remeasure(collector.Run("test"), dict(config.DEFAULTS))
    assert not called


def test_settings_and_manual_activation(conn):
    from trendradar.web import app

    client = app.test_client()
    assert client.post("/api/settings", json={"min_per_bron": 0}).status_code == 400
    assert budget.allocate([{"id": 1, "yield_score": 0.5, "priority": 1, "last_fetched": None}], 30, 0, 0)
    conn.execute("INSERT INTO sources(platform, kind, value, status, origin, added_at, yield_score, fetch_count) "
                 "VALUES ('tiktok', 'zoekterm', 'kroket', 'gepauzeerd', 'ontdekt', 'x', 0.05, 7)")
    conn.commit()
    sid = conn.execute("SELECT id FROM sources").fetchone()[0]
    assert client.patch(f"/api/sources/{sid}", json={"status": "actief"}).status_code == 200
    row = conn.execute("SELECT status, yield_score, fetch_count FROM sources WHERE id = ?", (sid,)).fetchone()
    assert (row["status"], row["yield_score"], row["fetch_count"]) == ("actief", None, 0)


def test_spend_is_reconciled_with_apify_account(conn, monkeypatch):
    from trendradar import collector

    ap = _apify_offline(monkeypatch)
    usage = iter([{"used_usd": 10.0, "limit_usd": 39.0, "cycle_end": "2026-11-15"},
                  {"used_usd": 12.5, "limit_usd": 39.0, "cycle_end": "2026-11-15"}])
    monkeypatch.setattr(ap, "account_usage", lambda: next(usage))
    collector._reconcile_account(collector.Run("test"))                       # eerste meetpunt
    budget.record(conn, "apify", 1.0, "tiktok", "ronde")                      # de tool telde $1
    conn.commit()
    collector._reconcile_account(collector.Run("test"))                       # Apify rekende $2.50
    assert budget.spent_usd(conn, "apify") == pytest.approx(2.5)


def test_other_websites_cannot_change_anything(conn):
    from trendradar.web import app

    client = app.test_client()
    body = {"post_id": "tiktok:x", "value": 1}
    assert client.post("/api/feedback", json=body).status_code == 200                 # geen Origin: curl, tests
    assert client.post("/api/feedback", json=body, headers={"Origin": "http://localhost"}).status_code == 200
    for headers in ({"Origin": "https://evil.example"}, {"Origin": "null"}, {"Sec-Fetch-Site": "cross-site"}):
        r = client.post("/api/run", headers=headers)
        assert r.status_code == 403 and "andere website" in r.get_json()["fout"], headers
    # Een vreemde domeinnaam die naar deze computer wijst (DNS-rebinding): ook lezen mag niet.
    assert client.get("/api/status", headers={"Host": "evil.example:8050"}).status_code == 403
    assert client.get("/api/status", headers={"Host": "127.0.0.1:8050"}).status_code == 200


def test_bad_input_gives_a_message_not_a_crash(conn):
    from trendradar.web import app

    client = app.test_client()
    for url in ("/api/posts?limit=veel", "/api/trends?min_rel=x", "/api/zaken?window=1j",
                "/api/supermarkt?days=abc", "/api/potentie?limit=tien", "/api/export/trends.csv?window=1j"):
        r = client.get(url)
        assert r.status_code == 400 and r.get_json()["fout"], url
    for path in ("/api/lijst", "/api/tips", "/api/feedback", "/api/sources", "/api/settings"):
        r = client.post(path, data="[1, 2]", content_type="application/json")
        assert r.status_code == 400 and r.get_json()["fout"], path
    assert client.post("/api/settings", json={"verdeling": [1]}).status_code == 400
    assert client.get("/api/bestaat-niet").get_json() == {"fout": "niet gevonden"}
    store_posts(conn, [post(i, f"Kroket nummer {i} in Zwolle", platform="tiktok") for i in range(3)], None)
    conn.commit()
    assert len(client.get("/api/posts?limit=-1").get_json()["posts"]) == 1           # niet 'alles'


def test_posts_by_id_and_partial_settings(conn):
    from trendradar.web import app

    client = app.test_client()
    store_posts(conn, [post(i, f"Kroket nummer {i} in Zwolle", platform="tiktok", hours_ago=24 * 20)
                       for i in range(4)], None)
    conn.commit()
    d = client.get("/api/posts?window=24u&ids=tiktok:1,tiktok:3").get_json()
    assert sorted(p["id"] for p in d["posts"]) == ["tiktok:1", "tiktok:3"] and d["total"] == 2
    assert client.post("/api/settings", json={"verdeling": {"tiktok": 0.5}}).status_code == 200
    assert client.get("/api/settings").get_json()["verdeling"] == {"instagram": 0.30, "tiktok": 0.5,
                                                                    "facebook": 0.25}
    assert client.post("/api/settings", json={"verdeling": {"tiktok": -1}}).status_code == 400


def test_csv_export_does_not_run_formulas(conn):
    from trendradar.web import app

    client = app.test_client()
    client.post("/api/lijst", json={"kind": "zaak", "ref": "x", "title": "=HYPERLINK(\"http://evil\")"})
    client.post("/api/lijst", json={"kind": "zaak", "ref": "y", "title": "@SUM(A1)", "note": "-2+3"})
    text = client.get("/api/export/lijst.csv").get_data(as_text=True)
    assert "'=HYPERLINK" in text and "'@SUM" in text and "'-2+3" in text
    assert "\n=HYPERLINK" not in text and ";=HYPERLINK" not in text


def test_frozen_supermarket_is_reported(conn, monkeypatch):
    from trendradar import collector
    from trendradar.sources import supermarket

    items = [{"id": "ah:1", "store": "Albert Heijn", "name": "Kroket", "size": "2 st", "price": 1.0, "url": None}]
    monkeypatch.setattr(supermarket, "fetch", lambda etag=None: ({"Albert Heijn": items}, "x"))
    settings = dict(config.DEFAULTS)
    collector.step_supermarkets(collector.Run("test"), settings, force=True)
    collector.step_supermarkets(collector.Run("test"), settings, force=True)
    assert db.get_state(conn, "supermarkt_status")["Albert Heijn"]["ok"]
    old = db.iso(db.now() - timedelta(days=collector.STALE_DAYS + 1))
    db.set_state(conn, "supermarkt_inhoud_Albert Heijn", {**db.get_state(conn, "supermarkt_inhoud_Albert Heijn"),
                                                          "since": old})
    conn.commit()
    collector.step_supermarkets(collector.Run("test"), settings, force=True)
    st = db.get_state(conn, "supermarkt_status")["Albert Heijn"]
    assert not st["ok"] and st["stil"] == old and "staat stil" in st["melding"]
    items.append({"id": "ah:2", "store": "Albert Heijn", "name": "Frikandel", "size": "", "price": 1.0, "url": None})
    collector.step_supermarkets(collector.Run("test"), settings, force=True)
    assert db.get_state(conn, "supermarkt_status")["Albert Heijn"]["ok"]


def test_backup_and_cleanup(conn, tmp_path):
    from trendradar import onderhoud

    store_posts(conn, [post("oud", "Kroket in Zwolle", platform="tiktok", hours_ago=24 * 400),
                       post("bewaard", "Frikandel in Zwolle", platform="tiktok", hours_ago=24 * 400),
                       post("nieuw", "Kroket in Ede", platform="tiktok", hours_ago=5)], None)
    conn.execute("INSERT INTO watchlist(kind, ref, title, added_at, updated_at) VALUES ('post', 'tiktok:bewaard', 'x', 'a', 'a')")
    conn.commit()
    settings = {**config.DEFAULTS, "backup_map": str(tmp_path / "bk")}
    gone = onderhoud.cleanup(conn, settings)
    conn.commit()
    ids = {r[0] for r in conn.execute("SELECT id FROM posts")}
    assert gone["posts"] == 1 and ids == {"tiktok:bewaard", "tiktok:nieuw"}
    assert not onderhoud.cleanup_due(conn)
    for _ in range(onderhoud.KEEP_BACKUPS + 2):
        onderhoud.make_backup(settings)  # zelfde minuut: zelfde naam, dus geen stapel kopieën
    files = list((tmp_path / "bk").glob("trendradar-*.db"))
    assert 1 <= len(files) <= onderhoud.KEEP_BACKUPS and not list((tmp_path / "bk").glob("*.deel"))
    import sqlite3
    assert sqlite3.connect(files[0]).execute("SELECT COUNT(*) FROM posts").fetchone()[0] == 2
    assert not onderhoud.backup_due(conn, settings)


def test_budget_shifts_to_the_platform_that_delivers(conn):
    rows = [("tiktok", "kroket", 10.0, 400), ("instagram", "snack", 10.0, 100), ("facebook", "friet", 1.0, 0)]
    for platform, value, usd, useful in rows:
        conn.execute("INSERT INTO sources(platform, kind, value, status, origin, added_at, cost_usd, useful_total) "
                     "VALUES (?, 'hashtag', ?, 'actief', 'start', 'x', ?, ?)", (platform, value, usd, useful))
    settings = dict(config.DEFAULTS)
    split = budget.smart_split(conn, settings)
    assert split["tiktok"] > settings["verdeling"]["tiktok"]          # 40 per dollar
    assert split["instagram"] < settings["verdeling"]["instagram"]    # 10 per dollar
    assert split["facebook"] == settings["verdeling"]["facebook"]     # nog te weinig uitgegeven om te oordelen
    assert split["instagram"] >= settings["verdeling"]["instagram"] / 2
    assert budget.smart_split(conn, {**settings, "verdeling_automatisch": False}) == settings["verdeling"]
    assert sum(budget.overview(conn, settings)["verdeling"].values()) in (99, 100, 101)


def test_trefkans_links_suggestions_to_own_posts(conn):
    from trendradar import trefkans

    old = db.iso(db.now() - timedelta(days=40))
    trefkans.log(conn, "kandidaat", "zin:kroket + kaassoufflé", "kroket + kaassoufflé",
                 trefkans._terms_words("kroket kaassoufflé"), ["Combinatie: kroket + kaassoufflé", "Doorbraak"])
    trefkans.log(conn, "potentie", "zaak:snackbardepoppe", "Snackbar de Poppe",
                 trefkans._terms_zaak("Snackbar de Poppe"), ["Eigen creatie: broodje lammert"])
    trefkans.log(conn, "potentie", "zaak:nooitgepost", "Nooit Gepost Bar", trefkans._terms_zaak("Nooit Gepost Bar"), [])
    conn.execute("UPDATE suggestions SET first_at = ?", (old,))
    day = (db.now() - timedelta(days=10)).strftime("%Y-%m-%d")
    for i, (cap, views) in enumerate([("Kroket met kaassoufflé erin!", 400000), ("Bij @snackbardepoppe gegeten", 300000),
                                      ("Gewone friet", 100000), ("Nog een friet", 100000)]):
        conn.execute("INSERT INTO own_posts(id, date, month, caption, views, followers_count, imported_at) "
                     "VALUES (?, ?, ?, ?, ?, 1, 'x')", (f"p{i}", day, day[:7], cap, views))
    r = trefkans.report(conn)
    assert r["suggested"] == 3 and r["posted"] == 2
    sig = {s["signal"]: s for s in r["signals"]}
    assert sig["combinatie"]["hits"] == 1 and sig["combinatie"]["lift"] == 2.0
    assert "eigen creatie" in sig
    assert trefkans.signal_of("12× views") == "veel views"


def test_own_posts_fetched_and_kept_on_import(conn, monkeypatch):
    from trendradar import collector, imports, ownposts

    conn.execute("INSERT INTO own_posts(id, date, month, caption, url, views, followers, followers_count, imported_at) "
                 "VALUES ('a', '2026-10-01', '2026-10', 'oud', 'https://www.instagram.com/reel/AAA111/', 1000, 50, 1, 'x')")
    conn.commit()
    real_normalize = apify.normalize
    ap = _apify_offline(monkeypatch)
    monkeypatch.setattr(ap, "normalize", real_normalize)  # hier juist de echte omzetting
    items = [{"id": "1", "shortCode": "AAA111", "url": "https://www.instagram.com/p/AAA111/", "caption": "oud",
              "ownerUsername": "snackspert", "timestamp": "2026-10-01T10:00:00.000Z", "videoPlayCount": 5000,
              "likesCount": 300, "commentsCount": 20},
             {"id": "2", "shortCode": "BBB222", "url": "https://www.instagram.com/p/BBB222/", "caption": "Nieuwe kroket!",
              "ownerUsername": "snackspert", "timestamp": "2026-10-20T10:00:00.000Z", "videoPlayCount": 8000,
              "likesCount": 400, "commentsCount": 30},
             {"id": "3", "shortCode": "CCC333", "url": "https://www.instagram.com/p/CCC333/", "caption": "collab",
              "ownerUsername": "iemandanders", "timestamp": "2026-10-21T10:00:00.000Z"}]
    monkeypatch.setattr(ap, "run_actor", lambda *a, **k: (items, 0.05))
    settings = {**config.DEFAULTS, "eigen_instagram": "@snackspert"}
    collector.step_own_posts(collector.Run("test"), settings)
    rows = {r["url"]: dict(r) for r in conn.execute("SELECT * FROM own_posts")}
    assert rows["https://www.instagram.com/reel/AAA111/"]["views"] == 5000          # bijgewerkt
    assert rows["https://www.instagram.com/reel/AAA111/"]["followers"] == 50        # export-cijfer blijft
    assert rows["https://www.instagram.com/p/BBB222/"]["note"] == "automatisch"
    assert len(rows) == 2                                                           # niet de post van een ander
    collector.step_own_posts(collector.Run("test"), settings)                       # pas over een week weer
    # Een nieuwe Excel-import zonder de automatische post laat die staan.
    imports.save(conn, {"posts": [{"id": "x", "url": "https://www.instagram.com/reel/AAA111/", "caption": "oud",
                                   "views": 6000, "month": "2026-10", "followers_count": 1}],
                        "spots": [], "routes": []})
    assert {r["url"] for r in conn.execute("SELECT url FROM own_posts")} == {
        "https://www.instagram.com/reel/AAA111/", "https://www.instagram.com/p/BBB222/"}
    assert ownposts.shortcode("https://www.instagram.com/snackspert/reel/XyZ-1/") == "XyZ-1"


def test_new_places_from_google_maps(conn, monkeypatch):
    from trendradar import collector, nieuwezaken

    ap = _apify_offline(monkeypatch)
    old = [{"placeId": f"old{i}", "title": f"Snackbar Oud {i}", "city": "Utrecht", "reviewsCount": 300,
            "totalScore": 4.1, "searchString": "snackbar Utrecht"} for i in range(3)]
    batches = iter([(old, 0.4), (old + [
        {"placeId": "n1", "title": "Frietkot De Nieuwe", "city": "Utrecht", "reviewsCount": 12, "totalScore": 4.8,
         "categoryName": "Snackbar", "searchString": "snackbar Utrecht"},
        {"placeId": "n2", "title": "Snackbar Gesloten", "city": "Utrecht", "reviewsCount": 3, "permanentlyClosed": True,
         "searchString": "snackbar Utrecht"},
        {"placeId": "n3", "title": "Al Lang Bekend", "city": "Utrecht", "reviewsCount": 900,
         "searchString": "snackbar Utrecht"}], 0.4)])
    calls = []
    monkeypatch.setattr(ap, "run_actor", lambda *a, **k: calls.append(a) or next(batches))
    settings = {**config.DEFAULTS, "maps_steden": ["Utrecht"], "maps_zoektermen": ["snackbar"]}
    collector.step_maps(collector.Run("test"), settings)
    assert calls[0][:2] == ("maps", "zoekterm") and calls[0][2] == ["snackbar Utrecht"]
    assert nieuwezaken.new_places(conn) == []                                   # nulmeting
    db.set_state(conn, "maps_laatst", {"at": db.iso(db.now() - timedelta(days=8)), "cities": [], "found": 0, "new": 0})
    conn.commit()
    collector.step_maps(collector.Run("test"), settings)
    names = [p["name"] for p in nieuwezaken.new_places(conn)]
    assert names == ["Frietkot De Nieuwe"]                                      # niet gesloten, niet al lang bekend
    assert nieuwezaken.new_places(conn)[0]["hot"]
    conn.execute("INSERT INTO spots(name, place, added_at, updated_at) VALUES ('Frietkot De Nieuwe', 'Utrecht', 'x', 'x')")
    assert nieuwezaken.new_places(conn) == []                                   # al in je tips
    assert budget.spent_usd(conn, "apify") == pytest.approx(0.8)


def test_supermarkets_via_apify(conn, monkeypatch):
    from trendradar import collector
    from trendradar.sources import supermarket

    ap = _apify_offline(monkeypatch)
    rounds = iter([
        [{"name": "AH Kroketten 4 st", "price": 2.49, "unit_size": "4 st", "url": "https://www.ah.nl/p/1"}],
        [{"name": "AH Kroketten 4 st", "price": 2.49, "unit_size": "4 st", "url": "https://www.ah.nl/p/1"},
         {"title": "AH Loaded Kipcorn Bites", "price": "€ 3,29", "url": "https://www.ah.nl/p/2"},
         {"name": "Jumbo Frikandel", "supermarket": "jumbo", "url": "https://www.jumbo.com/x"}],
    ])
    calls = []

    def fake(platform, kind, values, per_value, max_usd, **kw):
        calls.append((platform, kind, values, kw.get("extra_input")))
        return next(rounds), 0.3
    monkeypatch.setattr(ap, "run_actor", fake)
    cjb = {"Albert Heijn": [{"id": "ah:old", "store": "Albert Heijn", "name": "Oud", "size": "", "price": 1, "url": None}],
           "Spar": [{"id": "spar:1", "store": "Spar", "name": "Spar Kroket", "size": "", "price": 1, "url": None}]}
    monkeypatch.setattr(supermarket, "fetch", lambda etag=None: (cjb, "e"))
    settings = {**config.DEFAULTS, "supermarkt_apify_winkels": ["ah"]}
    collector.step_supermarkets(collector.Run("t"), settings, force=True)
    collector.step_supermarkets_apify(collector.Run("t"), settings, force=True)
    assert calls[0][:2] == ("supermarkt", "zoekterm") and calls[0][3]["supermarkets"] == ["ah"]
    assert calls[0][3]["maxResults"] == 12 * len(settings["supermarkt_zoektermen"])  # totaal, niet per woord
    status = db.get_state(conn, "supermarkt_status")
    assert status["Albert Heijn"]["bron"] == "apify" and "nulmeting" in status["Albert Heijn"]["melding"]
    assert status["Spar"]["bron"] == "cjb"
    assert not conn.execute("SELECT 1 FROM products WHERE id = 'ah:old'").fetchone()   # stilstaande bron niet gebruikt
    collector.step_supermarkets_apify(collector.Run("t"), settings, force=True)
    new = {r["name"]: r["price"] for r in conn.execute("SELECT name, price FROM products WHERE is_new = 1")}
    assert new == {"AH Loaded Kipcorn Bites": 3.29}                                    # Jumbo-resultaat niet bij AH
    assert budget.spent_usd(conn, "apify") == pytest.approx(0.6)
    assert supermarket.parse_apify({"productName": "X", "price": {"now": 1.5}}, "Plus", "plus")["price"] == 1.5


def test_supermarkets_via_apify_run_in_background(conn, monkeypatch):
    import threading
    from trendradar import collector

    _apify_offline(monkeypatch)
    done = threading.Event()
    monkeypatch.setattr(collector, "step_supermarkets_apify", lambda run, settings, force=False: done.wait(5))
    settings = {**config.DEFAULTS, "supermarkt_apify_winkels": ["ah"]}
    assert collector.start_supermarkets_apify(settings, True)
    assert not collector.start_supermarkets_apify(settings, True)        # niet twee tegelijk
    assert not collector.is_running()                                     # de gewone ronde wacht er niet op
    done.set()
    for _ in range(50):
        if collector._shop_lock.acquire(blocking=False):
            collector._shop_lock.release()
            break
        threading.Event().wait(0.05)
    row = conn.execute("SELECT trigger FROM runs ORDER BY id DESC LIMIT 1").fetchone()
    assert row["trigger"] == "supermarkt"
    assert collector.due(conn, settings)                                  # telt niet als gewone ronde


def test_supermarkets_on_fixed_times(conn, monkeypatch):
    from datetime import datetime, timezone

    from trendradar import collector

    _apify_offline(monkeypatch)
    settings = {**config.DEFAULTS, "supermarkt_apify_winkels": ["ah"]}
    tz = timezone.utc
    wed = datetime(2026, 10, 28, 12, 0, tzinfo=tz)                       # woensdag
    assert collector._last_slot(settings["supermarkt_tijden"], wed) == datetime(2026, 10, 26, 7, 0, tzinfo=tz)
    assert collector._last_slot(["do 17.00"], datetime(2026, 10, 29, 17, 0, tzinfo=tz)).day == 29
    assert collector._last_slot(["onzin", "maandag 25:00"], wed) is None
    assert collector.shop_due(conn, settings, wed)                       # nog nooit gedaan
    db.set_state(conn, "supermarkt_apify_laatst", "2026-10-26T07:05:00Z")
    assert not collector.shop_due(conn, settings, wed)                   # maandag al gedaan
    assert collector.shop_due(conn, settings, datetime(2026, 10, 29, 17, 1, tzinfo=tz))   # donderdag 17:00
    assert not collector.shop_due(conn, {**settings, "supermarkt_apify_aan": False}, wed.replace(day=29, hour=18))


def test_budget_page_with_spend_without_platform(conn):
    from trendradar.web import app

    budget.record(conn, "apify", 0.5, None, "Google Maps")
    budget.record(conn, "apify", 0.2, "tiktok", "ronde")
    conn.commit()
    r = app.test_client().get("/api/budget")
    assert r.status_code == 200
    assert r.get_json()["data_per_platform_eur"]["overig"] > 0
