"""Excel-bestanden inlezen: je postoverzicht en je tips en routes.

De tool herkent zelf wat er in een bestand staat, aan de kolomnamen:

- een tabblad met 'Weergaven' en 'Caption'  -> je eigen posts
- een tabblad met 'Route' en 'Stop 1'       -> routes
- een tabblad met 'Naam' (of 'Naam zaak')   -> tips (spots)

Tabbladen met samenvattingen (Per maand, Samenvatting) slaat hij over: die
rekent de tool zelf uit. Opnieuw importeren kan altijd. Je posts en routes
worden dan vervangen; bij tips blijven jouw status en notities bewaard.
"""
from __future__ import annotations

import io
import json
import re
from datetime import date, datetime, timedelta

from . import db, taste

# Kolomnamen, Nederlands en Engels (een export uit Instagram of Meta Business Suite heet anders).
POST_COLUMNS = {
    "maand": "month", "month": "month",
    "datum": "date", "date": "date", "gepost op": "date", "publicatiedatum": "date", "publish time": "date",
    "type": "type", "soort post": "type", "post type": "type",
    "caption (begin)": "caption", "caption": "caption", "tekst": "caption", "description": "caption",
    "link": "url", "url": "url", "permalink": "url",
    "weergaven": "views", "views": "views", "plays": "views", "afspeelacties": "views",
    "bereik": "reach", "reach": "reach",
    "likes": "likes", "vind-ik-leuks": "likes",
    "reacties": "comments", "comments": "comments",
    "gedeeld": "shares", "shares": "shares",
    "opgeslagen": "saves", "saves": "saves",
    "volgers": "followers", "follows": "followers", "followers": "followers", "nieuwe volgers": "followers",
    "format": "format", "kanaal": "channel", "locatie": "location", "herkenbaarheid": "recognition",
    "indeling": "classified", "volgers telt mee": "followers_count", "notitie": "note",
}
NUMBERS = {"views", "reach", "likes", "comments", "shares", "saves", "followers"}
_VIEW_KEYS = {k for k, v in POST_COLUMNS.items() if v == "views"}
_TEXT_KEYS = {k for k, v in POST_COLUMNS.items() if v in ("caption", "url")}
_NAME_KEYS = ("naam", "naam zaak", "name", "zaak")
_TOTAL_WORDS = ("totaal", "gemiddelde", "gemiddeld", "mediaan", "som", "total", "average", "median")
_MONTHS = {"jan": 1, "januari": 1, "january": 1, "feb": 2, "februari": 2, "february": 2, "mrt": 3, "maart": 3,
           "mar": 3, "march": 3, "apr": 4, "april": 4, "mei": 5, "may": 5, "jun": 6, "juni": 6, "june": 6,
           "jul": 7, "juli": 7, "july": 7, "aug": 8, "augustus": 8, "august": 8, "sep": 9, "sept": 9,
           "september": 9, "okt": 10, "oktober": 10, "oct": 10, "october": 10, "nov": 11, "november": 11,
           "dec": 12, "december": 12}


class ImportError_(ValueError):
    """Het bestand bevat niets wat de tool herkent."""


def _key(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _number(v):
    """12345, '1.234', '12,5K', '1,2M', '1234.0' -> een geheel getal; anders None."""
    if v is None or v == "" or isinstance(v, (datetime, date)):
        return None
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)):
        return int(round(v))
    s = str(v).strip().lower().replace("\u00a0", "").replace(" ", "")
    m = re.fullmatch(r"([0-9][0-9.,]*)(k|m|mln|mio|duizend)?", s)
    if not m:
        digits = re.sub(r"[^\d]", "", s)
        return int(digits) if digits else None
    num, unit = m.groups()
    mult = {"k": 1e3, "duizend": 1e3, "m": 1e6, "mln": 1e6, "mio": 1e6}.get(unit, 1)
    if "," in num and "." in num:
        # Het laatste teken is de decimaal: '1.234,5' of '1,234.5'.
        num = num.replace(".", "").replace(",", ".") if num.rfind(",") > num.rfind(".") else num.replace(",", "")
    elif "," in num or "." in num:
        sep = "," if "," in num else "."
        parts = num.split(sep)
        decimal = unit or (len(parts) == 2 and len(parts[1]) != 3)
        num = num.replace(sep, ".") if decimal else num.replace(sep, "")
    try:
        return int(round(float(num) * mult))
    except ValueError:
        return None


def _to_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 30000 < v < 70000:  # Excel-datumgetal
        return date(1899, 12, 30) + timedelta(days=int(v))
    s = _text(v)
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(s[:10] if fmt.startswith("%Y") else s, fmt).date()
        except ValueError:
            continue
    return None


def _to_month(v) -> str | None:
    """'2026-03', datumcel, '17-03-2026', 'maart 2026', 'mrt-26', '03-2026' -> '2026-03'."""
    d = _to_date(v)
    if d:
        return f"{d.year:04d}-{d.month:02d}"
    s = _key(v)
    if m := re.fullmatch(r"(\d{4})[-/.](\d{1,2})", s):
        y, mo = int(m.group(1)), int(m.group(2))
    elif m := re.fullmatch(r"(\d{1,2})[-/.](\d{4})", s):
        y, mo = int(m.group(2)), int(m.group(1))
    elif m := re.fullmatch(r"([a-z]+)\.?[\s\-'’]*(\d{2}|\d{4})", s):
        mo = _MONTHS.get(m.group(1))
        y = int(m.group(2))
        y = y + 2000 if y < 100 else y
        if not mo:
            return None
    else:
        return None
    return f"{y:04d}-{mo:02d}" if 1 <= mo <= 12 and 2000 <= y <= 2100 else None


def _header(ws, max_scan: int = 10):
    """(rijnummer, [kolomnamen], soort) van de eerste rij met herkenbare kolomnamen.

    Een titelregel als 'Overzicht 2026 | bijgewerkt op ...' erboven is geen probleem.
    """
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan), start=1):
        names = [_key(c.value) for c in row]
        if any(n in _VIEW_KEYS for n in names) and any(n in _TEXT_KEYS for n in names):
            return i, names, "posts"
        if "route" in names and "stop 1" in names:
            return i, names, "routes"
        if any(n in _NAME_KEYS for n in names):
            return i, names, "spots"
    return None, [], None


def _list_name(title: str) -> str:
    t = title.strip()
    low = t.lower()
    if "buiten" in low and ("a'dam" in low or "amsterdam" in low):
        return "Buiten Amsterdam"
    if low.startswith("alle spots"):
        return "Spots"
    return t


def read_workbook(data: bytes) -> dict:
    """Leest een .xlsx en geeft {'posts': [...], 'spots': [...], 'routes': [...]}."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # zip- of xml-fout: geen geldig Excel-bestand
        raise ImportError_("Dit is geen Excel-bestand (.xlsx) dat de tool kan lezen.") from exc
    out = {"posts": [], "spots": [], "routes": []}
    for ws in wb.worksheets:
        row_no, names, kind = _header(ws)
        if kind == "posts":
            out["posts"] += _read_posts(ws, row_no, names)
        elif kind == "routes":
            out["routes"] += _read_routes(ws, row_no, names)
        elif kind == "spots":
            out["spots"] += _read_spots(ws, row_no, names)
    # Tabbladen als 'Samenwerkingen' herhalen posts uit 'Posts'. Houd per post de vollediger rij.
    best: dict[str, dict] = {}
    for p in out["posts"]:
        filled = sum(1 for v in p.values() if v not in (None, ""))
        if p["id"] not in best or filled > sum(1 for v in best[p["id"]].values() if v not in (None, "")):
            best[p["id"]] = p
    out["posts"] = list(best.values())
    if not any(out.values()):
        raise ImportError_("Geen posts, tips of routes gevonden. Herkenbare kolommen zijn bijvoorbeeld "
                           "'Weergaven' en 'Caption' (posts), 'Naam' (tips) of 'Route' en 'Stop 1' (routes).")
    return out


def _read_posts(ws, row_no, names) -> list[dict]:
    cols = {i: POST_COLUMNS[n] for i, n in enumerate(names) if n in POST_COLUMNS}
    rows = []
    for row in ws.iter_rows(min_row=row_no + 1):
        p, raw = {}, {}
        for i, field in cols.items():
            cell = row[i] if i < len(row) else None
            if cell is None:
                continue
            raw[field] = cell.value
            if field == "url":
                p["url"] = cell.hyperlink.target if cell.hyperlink else (
                    _text(cell.value) if _text(cell.value).startswith("http") else "")
            elif field in NUMBERS:
                p[field] = _number(cell.value)
            elif field not in ("month", "date"):
                p[field] = _text(cell.value)
        # Geen tekst en geen link: een lege regel of een totaalregel, geen post.
        if not p.get("caption") and not p.get("url"):
            continue
        # Een totaalregel ('Totaal', 'Gemiddelde:'), niet een caption die met 'Sommige...' begint.
        if not p.get("url") and re.match(r"^(%s)\b" % "|".join(_TOTAL_WORDS), _key(p.get("caption"))):
            continue
        d = _to_date(raw.get("date")) if raw.get("date") not in (None, "") else None
        p["date"] = d.isoformat() if d else None
        p["month"] = _to_month(raw.get("month")) if raw.get("month") not in (None, "") else None
        if not p["month"] and d:
            p["month"] = f"{d.year:04d}-{d.month:02d}"
        p["followers_count"] = 0 if _key(p.get("followers_count")) == "nee" else 1
        p["id"] = p.get("url") or f"{p.get('month') or ''}:{(p.get('caption') or '')[:60]}"
        rows.append(p)
    return rows


def _read_spots(ws, row_no, names) -> list[dict]:
    def col(*options):
        return next((i for i, n in enumerate(names) if n in options), None)
    c_name, c_place = col(*_NAME_KEYS), col("locatie", "plaats", "stad", "place", "city")
    c_spec = col("specialiteit", "beschrijving", "opmerking", "opmerkingen", "specialty")
    c_region = col("regio", "provincie", "region")
    lst = _list_name(ws.title)
    title = ws.title.lower()
    default_place = "Amsterdam" if c_place is None and "amsterdam" in title and "buiten" not in title else ""
    out = []
    for row in ws.iter_rows(min_row=row_no + 1, values_only=True):
        name = _text(row[c_name]) if c_name is not None and c_name < len(row) else ""
        if not name or re.match(r"^(%s)\b" % "|".join(_TOTAL_WORDS), _key(name)):
            continue
        get = lambda c: _text(row[c]) if c is not None and c < len(row) else ""  # noqa: E731
        place = get(c_place) or default_place
        if place.upper() == "ONBEKEND":
            place = ""
        region = get(c_region)
        if region.upper() == "ONBEKEND":
            region = ""
        if not region and place == "Amsterdam":
            region = "Noord-Holland"
        out.append({"name": name, "place": place, "specialty": get(c_spec), "region": region, "list": lst})
    return out


def _read_routes(ws, row_no, names) -> list[dict]:
    stops = [i for i, n in enumerate(names) if n.startswith("stop")]
    def col(prefix):
        return next((i for i, n in enumerate(names) if n.startswith(prefix)), None)
    c_name, c_region, c_km, c_time, c_note = col("route"), col("regio"), col("km"), col("rijtijd"), col("opmerking")
    out = []
    for row in ws.iter_rows(min_row=row_no + 1, values_only=True):
        get = lambda c: _text(row[c]) if c is not None and c < len(row) else ""  # noqa: E731
        name = get(c_name)
        route_stops = [get(i) for i in stops if get(i)]
        if not name or not route_stops:
            continue
        out.append({"name": name, "region": get(c_region), "stops": route_stops, "km": get(c_km),
                    "duration": get(c_time), "note": get(c_note)})
    return out


def parse_stop(stop: str) -> dict:
    """'Snackbar de Poppe (Coevorden) - broodje lammert' -> naam, plaats, specialiteit.

    Ook met haakjes in de plaats ('Paul van Gurp (Colmschate (Deventer))') of in de
    specialiteit ('broodje lammert (hamburger met kaassoufflé)').
    """
    text = stop.strip()
    open_at = text.find("(")
    if open_at <= 0:
        return {"name": text, "place": "", "specialty": ""}
    depth, close_at = 0, -1
    for k in range(open_at, len(text)):
        depth += {"(": 1, ")": -1}.get(text[k], 0)
        if depth == 0:
            close_at = k
            break
    if close_at < 0:
        return {"name": text, "place": "", "specialty": ""}
    rest = text[close_at + 1:].strip()
    if rest and not rest.startswith(("-", "–")):
        return {"name": text, "place": "", "specialty": ""}
    return {"name": text[:open_at].strip(), "place": text[open_at + 1:close_at].strip(),
            "specialty": rest.lstrip("-–").strip()}


def _place_key(place: str) -> str:
    """'Colmschate (Deventer)' en 'Colmschate' horen bij elkaar."""
    return re.sub(r"\s*\(.*?\)", "", place or "").split(",")[0].strip().lower()


def spot_key(name: str, place: str) -> tuple[str, str]:
    return taste.zaak_key(name), _place_key(place)


_STATUS_ORDER = ["te proberen", "gepland", "bezocht", "gereviewd", "afgevallen"]


def merge_duplicate_spots(conn) -> int:
    """Voeg dubbele tips samen ('Snackbar De Poppe' en 'Snackbar de Poppe', 'Colmschate' en
    'Colmschate (Deventer)'): lijsten en notities samen, de verste status blijft."""
    groups: dict[tuple, list[dict]] = {}
    for r in conn.execute("SELECT * FROM spots ORDER BY id"):
        d = db.row_to_dict(r)
        groups.setdefault(spot_key(d["name"], d["place"]), []).append(d)
    merged = 0
    for rows in groups.values():
        if len(rows) < 2 or not rows[0]["name"]:
            continue
        keep, rest = rows[0], rows[1:]
        lists = list(dict.fromkeys(x for r in rows for x in r["lists"]))
        notes = " · ".join(dict.fromkeys(r["note"] for r in rows if r["note"]))
        spec = max((r["specialty"] or "" for r in rows), key=len)
        status = max((r["status"] for r in rows),
                     key=lambda st: _STATUS_ORDER.index(st) if st in _STATUS_ORDER else 0)
        place = max((r["place"] or "" for r in rows), key=len)
        conn.execute("UPDATE spots SET lists = ?, note = ?, specialty = ?, status = ?, place = ? WHERE id = ?",
                     (json.dumps(lists), notes, spec, status, place, keep["id"]))
        conn.execute(f"DELETE FROM spots WHERE id IN ({','.join('?' for _ in rest)})", [r["id"] for r in rest])
        merged += len(rest)
    return merged


def save(conn, parsed: dict) -> dict:
    """Slaat het resultaat van read_workbook op. Geeft per soort het aantal terug."""
    now = db.iso(db.now())
    result = {"posts": 0, "spots": 0, "spots_new": 0, "routes": 0}
    if parsed["posts"]:
        # Het Excel-overzicht is leidend: alles vervangen voorkomt dubbele posts.
        conn.execute("DELETE FROM own_posts")
        fields = ["id", "date", "month", "type", "caption", "url", "views", "reach", "likes", "comments",
                  "shares", "saves", "followers", "format", "channel", "location", "recognition",
                  "classified", "followers_count", "note"]
        for p in parsed["posts"]:
            conn.execute(
                f"INSERT OR REPLACE INTO own_posts({', '.join(fields)}, imported_at) "
                f"VALUES ({', '.join('?' for _ in fields)}, ?)",
                [p.get(f) if p.get(f) != "" else None for f in fields] + [now])
        result["posts"] = len(parsed["posts"])
    if parsed["spots"]:
        existing = {}
        for r in conn.execute("SELECT id, name, place, lists, specialty, region FROM spots"):
            existing.setdefault(spot_key(r["name"], r["place"]), dict(r))
        for s in parsed["spots"]:
            key = spot_key(s["name"], s["place"])
            old = existing.get(key)
            if old:
                lists = json.loads(old["lists"] or "[]")
                if s["list"] not in lists:
                    lists.append(s["list"])
                spec = s["specialty"] if len(s["specialty"]) > len(old["specialty"] or "") else old["specialty"]
                conn.execute("UPDATE spots SET lists = ?, specialty = ?, region = ?, updated_at = ? WHERE id = ?",
                             (json.dumps(lists), spec, s["region"] or old["region"], now, old["id"]))
                old.update(lists=json.dumps(lists), specialty=spec)
            else:
                cur = conn.execute("INSERT INTO spots(name, place, specialty, region, lists, added_at, updated_at) "
                                   "VALUES (?, ?, ?, ?, ?, ?, ?) ON CONFLICT(name, place) DO NOTHING",
                                   (s["name"], s["place"], s["specialty"], s["region"], json.dumps([s["list"]]),
                                    now, now))
                if cur.rowcount:
                    existing[key] = {"id": cur.lastrowid, "name": s["name"], "place": s["place"],
                                     "lists": json.dumps([s["list"]]), "specialty": s["specialty"],
                                     "region": s["region"]}
                    result["spots_new"] += 1
            result["spots"] += 1
        merge_duplicate_spots(conn)
    if parsed["routes"]:
        # Bijwerken op naam, zodat je status ('bezocht') blijft staan.
        old = {r["name"]: r["id"] for r in conn.execute("SELECT id, name FROM routes WHERE origin = 'import'")}
        seen = set()
        for r in parsed["routes"]:
            stops = json.dumps(r["stops"], ensure_ascii=False)
            if r["name"] in old:
                conn.execute("UPDATE routes SET region = ?, stops = ?, km = ?, duration = ?, note = ? WHERE id = ?",
                             (r["region"], stops, r["km"], r["duration"], r["note"], old[r["name"]]))
            else:
                conn.execute("INSERT INTO routes(name, region, stops, km, duration, note, added_at) "
                             "VALUES (?, ?, ?, ?, ?, ?, ?)",
                             (r["name"], r["region"], stops, r["km"], r["duration"], r["note"], now))
            seen.add(r["name"])
        gone = [rid for name, rid in old.items() if name not in seen]
        if gone:
            conn.execute(f"DELETE FROM routes WHERE id IN ({','.join('?' for _ in gone)})", gone)
        result["routes"] = len(parsed["routes"])
    return result
